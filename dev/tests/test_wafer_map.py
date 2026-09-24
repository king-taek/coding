"""Wafer map — 평면 좌표 변환(순수) · 뷰/시트(헤드리스) · 엑셀 시트.

지키는 계약:

- **Camtek 역변환은 원시 stage 좌표를 그대로 되돌린다** — ``DefectCoord``(die 인덱스 +
  die 내부)를 평면으로 놓은 값이 ``(X − Center_X, Center_Y − Y)`` 와 floor 오차(1 µm)
  안에서 같다.  KLA 도 마찬가지로 ``(XINDEX·px + XREL − cx, YINDEX·py + YREL − cy)``.
  이게 틀리면 die 격자와 점이 어긋난다.
- 중심이 없으면 ``가정`` 등급으로 표시하고 점은 원 안에 놓인다(±½ pitch).
- 좌표를 못 놓은 사진은 조용히 사라지지 않고 ``unplaced`` 에 남는다.
- LIVE 파일명 사진만 든 폴더(INI 항목 없음)도 찍힌다 — 같은 결함의 INI 경로와 같은 자리.
  pitch 가 파일에 없으면 상수를 **가정**으로 쓰고 범례에 표기, 가정이 반증되면 안 찍는다.
- 격자선은 원 안만, 개수는 지름/pitch 근처 — die 8만 개도 선 몇백 개다.
- 뷰: 휠은 커서 기준 확대, 점 판정은 HIT_PX 이내, PNG 렌더는 화면과 같은 함수.
- 결과 시트는 기준/검증 두 맵, 셋업 시트는 폴더 안내 → 슬롯 폴더면 맵 하나, LOT 폴더면
  전체 합산 + '슬롯 선택…' 으로 일부만.  맵은 워커가 만들고 썸네일을 미리 굽는다.
- 점 더블클릭이 ``point_activated`` 를 낸다(단일 클릭은 아무것도 열지 않는다).
- die 색칠: 결함이 든 **그 칸만**, **결함 점과 같은 색**으로 칠하고 점은 생략한다.
  토글은 **셋업 단계 화면에만** 있다(결과 단계는 점 색이 곧 매치됨/미매치).
  칠하기 모드에서는 칸 아무 데나 집어도 그 결함이 잡힌다.
- 노치 회전: 90° 단위 시계 방향.  점·격자·노치·히트 판정이 **같이** 돈다(한 곳에서만
  회전하므로).  평면 좌표(``MapPoint.x/y``)는 안 바뀐다 — 보기 상태일 뿐이다.
- 사진 1장 선택: 고른 사진 하나만 찍고, 원·격자는 그 사진 폴더의 기하 그대로다.
- 전체화면: 주변 표시(설명·버튼줄·제목·범례)를 감춰 맵이 실제로 **커진다**.  나가기는
  ESC 와 떠 있는 버튼 둘 다.  닫을 때 내가 키운 창은 되돌린다.  안 보이는 창은 건드리지
  않는다(뜰 생각 없던 창을 띄우지 않는다).
- 엑셀: 렌더러가 주입되고 ``slot_images`` 가 있으면 'Wafer Map' 시트에 슬롯 행 +
  LOT 합산 행이 생긴다.  워커는 ui 를 import 하지 않는다(렌더러는 인자).
"""

from __future__ import annotations

import math
import os
from pathlib import Path

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from aoi_verification.app.coords import wafer_map as wm          # noqa: E402
from aoi_verification.app.coords import (camtek_ini, kla_info,   # noqa: E402
                                         resolve_batch, wafer_geometry)

PX, PY = 37247.7, 44905.4
DIA = 300000.0
CX, CY = 165994.0, 202629.0        # T254 실측 중심(§6-L) — 값 자체는 임의여도 된다


@pytest.fixture(autouse=True)
def _clear_caches():
    for fn in (wm._camtek, wm._kla, camtek_ini.load_folder, camtek_ini.load_raw_folder,
               camtek_ini.load_abs_folder, kla_info.load_folder,
               wafer_geometry.camtek_geometry, wafer_geometry.kla_geometry,
               wafer_geometry.live_geometry):
        fn.cache_clear()
    yield


def _camtek_folder(tmp_path: Path, entries, *, center: bool = True) -> Path:
    """entries: [(stem, X, Y)] — Col/Row 는 pitch 검산이 통과하도록 좌표에서 만든다."""
    folder = tmp_path / "camtek"
    folder.mkdir(parents=True)
    ini = "\n".join(
        f"[{stem}.jpeg]\nX={X}\nY={Y}\nCol={math.floor(X / PX)}\nRow={math.floor(Y / PY)}\n"
        for stem, X, Y in entries)
    (folder / "ColorImageGrabingInfo.ini").write_text(ini, encoding="utf-8")
    geom = (f"[Geometry]\nDieStep_X={PX:.6f}\nDieStep_Y={PY:.6f}\n"
            f"[Geometric]\nDiameter={DIA:.6f}\n")
    if center:
        geom += f"Center_X={CX:.6f}\nCenter_Y={CY:.6f}\n"
    (folder / "Params_WaferInfo.ini").write_text(geom, encoding="utf-8")
    for stem, _, _ in entries:
        (folder / f"{stem}.jpeg").write_bytes(b"")
    return folder


def _write_die_map(folder: Path, cells, *, repeat: int = 1) -> None:
    """``s_DieLocation.dat`` + 사이드카 — cells: {(x_index, y_index)} stage 인덱스."""
    import struct
    recs = b"".join(bytes(16) + struct.pack("<dd", i * PX + 10.0, j * PY + 10.0)
                    for i, j in sorted(cells)) * repeat
    (folder / wafer_geometry._DIE_MAP_FILE).write_bytes(recs)
    (folder / (wafer_geometry._DIE_MAP_FILE + ".md")).write_text(
        '<root><RecordSize Size="32"/><Fields>'
        '<Field Name="x" Id="1" Offset="16" Vartype="5"/>'
        '<Field Name="y" Id="2" Offset="24" Vartype="5"/>'
        '</Fields></root>', encoding="utf-8")


def _kla_folder(tmp_path: Path, defects, *, center: bool = True) -> Path:
    """defects: [(stem, XREL, YREL, XINDEX, YINDEX)]."""
    folder = tmp_path / "kla"
    folder.mkdir(parents=True)
    head = (f"FileVersion 1 2;\nDiePitch {PX:.6e} {PY:.6e};\n"
            f"SampleSize 1 300;\nWaferID \"W1\";\n")
    if center:
        head += f"SampleCenterLocation {CX:.6e} {CY:.6e};\n"
    body = "".join(
        f"TiffFileName {stem}.jpg\n 1 100.0 200.0 {xr} {yr} {xi} {yi} 0\n"
        for stem, xr, yr, xi, yi in defects)
    (folder / "W1.001").write_text(head + body, encoding="utf-8")
    for stem, *_ in defects:
        (folder / f"{stem}.jpg").write_bytes(b"")
    return folder


# ---------------------------------------------------------------------------
# 순수 변환
# ---------------------------------------------------------------------------
class TestCamtekPlane:
    def test_roundtrip_to_stage_frame(self, tmp_path):
        """die-local → 평면 = (X − Cx, Cy − Y).  floor 오차 1 µm 이내."""
        entries = [("a", 150000.0, 210000.0), ("b", 60000.5, 120000.25),
                   ("c", 250000.0, 300000.0)]
        folder = _camtek_folder(tmp_path, entries)
        coords = resolve_batch([folder / f"{s}.jpeg" for s, _, _ in entries])
        data = wm.build_map(coords)
        assert data.frame is not None and data.frame.center_source == wm.SOURCE_OBSERVED
        assert not data.unplaced
        by_name = {p.path.stem: p for p in data.points}
        for stem, X, Y in entries:
            p = by_name[stem]
            assert p.x == pytest.approx(X - CX, abs=1.0)
            assert p.y == pytest.approx(CY - Y, abs=1.0)
            assert p.col is not None and p.row is not None

    def test_grid_phase_matches_die_boundaries(self, tmp_path):
        """격자 경계는 stage 의 k·pitch 자리 — 평면에서 (k·PX − CX, CY − k·PY)."""
        folder = _camtek_folder(tmp_path, [("a", 150000.0, 210000.0)])
        frame = wm.frame_for_folder(folder, "camtek")
        xs, ys = wm.grid_lines(frame)
        assert xs and ys
        for x in xs:
            assert ((x + CX) / PX) == pytest.approx(round((x + CX) / PX), abs=1e-6)
            assert abs(x) <= DIA / 2
        for y in ys:
            assert ((CY - y) / PY) == pytest.approx(round((CY - y) / PY), abs=1e-6)
            assert abs(y) <= DIA / 2
        # 선 개수 ≈ 지름/pitch (+1) — die 8만 개(pitch ~1 mm)도 몇백 개다.
        assert abs(len(xs) - DIA / PX) <= 2
        assert abs(len(ys) - DIA / PY) <= 2

    def test_die_grid_segments_cover_exactly_the_whole_dies(self, tmp_path):
        """격자 선분은 온전한 die(네 꼭짓점이 원 안)의 변을 전부, 그리고 그것만 덮는다."""
        folder = _camtek_folder(tmp_path, [("a", 150000.0, 210000.0)])
        frame = wm.frame_for_folder(folder, "camtek")
        xs, ys = wm.grid_lines(frame)
        r = frame.radius
        inside = lambda x, y: x * x + y * y <= r * r + 1e-6       # noqa: E731
        want = set()        # 단위 변: ("v", i, j) = x=xs[i] 위 ys[j]~ys[j+1]
        for i in range(len(xs) - 1):
            for j in range(len(ys) - 1):
                if all(inside(x, y) for x in xs[i:i + 2] for y in ys[j:j + 2]):
                    want |= {("v", i, j), ("v", i + 1, j),
                             ("h", i, j), ("h", i, j + 1)}
        assert want
        got = set()
        for x1, y1, x2, y2 in wm.die_grid_segments(frame):
            assert inside(x1, y1) and inside(x2, y2)
            if x1 == x2:
                i = xs.index(x1)
                got |= {("v", i, j) for j in range(ys.index(y1), ys.index(y2))}
            else:
                j = ys.index(y1)
                got |= {("h", i, j) for i in range(xs.index(x1), xs.index(x2))}
        assert got == want

    def test_die_map_cells_are_drawn_as_is(self, tmp_path):
        """장비 die 맵이 채택되면 격자는 **그 칸만** — 계산(원 안 판정)으로 덮어쓰지 않는다.
        맵에서 가장자리 die 하나를 빼 두면(계산으로는 '있는' 칸) 그 칸은 그려지지 않는다."""
        folder = _camtek_folder(tmp_path, [("a", 150000.0, 210000.0)])
        r2 = (DIA / 2) ** 2
        full = {(i, j) for i in range(10) for j in range(10)
                if all((x - CX) ** 2 + (y - CY) ** 2 <= r2
                       for x in (i * PX, (i + 1) * PX) for y in (j * PY, (j + 1) * PY))}
        top = min(j for _, j in full)
        dropped = max(c for c in full if c[1] == top)        # 맨 윗줄 오른쪽 끝 die
        cells = full - {dropped}
        _write_die_map(folder, cells, repeat=3)     # 레코드 ≥ _MIN_DIE_MAP
        frame = wm.frame_for_folder(folder, "camtek")
        assert frame.die_cells == {(i, -(j + 1)) for i, j in cells}

        want = set()
        for i, k in frame.die_cells:
            want |= {("v", i, k), ("v", i + 1, k), ("h", i, k), ("h", i, k + 1)}
        kx = lambda x: round((x - frame.grid_x0) / PX)          # noqa: E731
        ky = lambda y: round((y - frame.grid_y0) / PY)          # noqa: E731
        got = set()
        for x1, y1, x2, y2 in wm.die_grid_segments(frame):
            if x1 == x2:
                got |= {("v", kx(x1), k) for k in range(ky(y1), ky(y2))}
            else:
                got |= {("h", i, ky(y1)) for i in range(kx(x1), kx(x2))}
        assert got == want
        # 뺀 die 의 바깥쪽 두 변(위·오른쪽)은 어느 이웃과도 공유되지 않는다.
        di, dk = dropped[0], -(dropped[1] + 1)
        assert ("h", di, dk + 1) not in got and ("v", di + 1, dk) not in got

    def test_partial_die_map_falls_back_to_computed_and_says_so(self, tmp_path):
        """부분 맵(유도값과 ±1 밖)은 기하가 버린다 → 격자도 계산으로, 범례에 그 사실을."""
        pytest.importorskip("PyQt6.QtWidgets")
        from aoi_verification.app import i18n
        from aoi_verification.app.ui.widgets.wafer_map_dialog import _MapPanel
        folder = _camtek_folder(tmp_path, [("a", 150000.0, 210000.0)])
        _write_die_map(folder, {(4 + n % 2, 3 + n % 2) for n in range(2)}, repeat=30)
        data = wm.build_map(resolve_batch([folder / "a.jpeg"]))
        assert data.frame.die_cells is None
        assert i18n.KO.WAFER_MAP_GRID_COMPUTED in _MapPanel.legend_text(data)

        wm._camtek.cache_clear()
        wafer_geometry.camtek_geometry.cache_clear()
        r2 = (DIA / 2) ** 2                                     # 온전한 맵 → 표기 없음
        full = {(i, j) for i in range(10) for j in range(10)
                if all((x - CX) ** 2 + (y - CY) ** 2 <= r2
                       for x in (i * PX, (i + 1) * PX) for y in (j * PY, (j + 1) * PY))}
        _write_die_map(folder, full, repeat=3)
        data = wm.build_map(resolve_batch([folder / "a.jpeg"]))
        assert data.frame.die_cells
        assert i18n.KO.WAFER_MAP_GRID_COMPUTED not in _MapPanel.legend_text(data)

    def test_no_center_is_assumed_and_inside_wafer(self, tmp_path):
        folder = _camtek_folder(tmp_path, [("a", 150000.0, 210000.0)], center=False)
        data = wm.build_map(resolve_batch([folder / "a.jpeg"]))
        assert data.frame is not None
        assert data.frame.center_source == wm.SOURCE_ASSUMED
        (p,) = data.points
        assert math.hypot(p.x, p.y) <= DIA / 2

    def test_unresolvable_goes_to_unplaced(self, tmp_path):
        folder = _camtek_folder(tmp_path, [("a", 150000.0, 210000.0)])
        stray = folder / "no_such_entry.jpeg"
        stray.write_bytes(b"")
        data = wm.build_map(resolve_batch([folder / "a.jpeg", stray]))
        assert len(data.points) == 1
        assert data.unplaced == (stray,)

    def test_matched_flag(self, tmp_path):
        folder = _camtek_folder(tmp_path, [("a", 150000.0, 210000.0),
                                           ("b", 60000.0, 120000.0)])
        paths = [folder / "a.jpeg", folder / "b.jpeg"]
        data = wm.build_map(resolve_batch(paths), matched={paths[0]})
        flags = {p.path.stem: p.matched for p in data.points}
        assert flags == {"a": True, "b": False}
        neutral = wm.build_map(resolve_batch(paths))
        assert all(p.matched is None for p in neutral.points)


def _live_folder(tmp_path: Path, stems, *, params: bool = True) -> Path:
    """LIVE 파일명 사진만 든 폴더 — ColorImageGrabingInfo.ini 없음."""
    folder = tmp_path / "live"
    folder.mkdir(parents=True)
    if params:
        (folder / "Params_WaferInfo.ini").write_text(
            f"[Geometry]\nDieStep_X={PX:.6f}\nDieStep_Y={PY:.6f}\n"
            f"[Geometric]\nDiameter={DIA:.6f}\nCenter_X={CX:.6f}\nCenter_Y={CY:.6f}\n",
            encoding="utf-8")
    for stem in stems:
        (folder / f"{stem}.jpg").write_bytes(b"")
    return folder


class TestLivePlane:
    """★ LIVE 사진만 든 폴더는 예전엔 기하가 None 이라 전부 '좌표 없음' 이었다."""

    def test_live_only_folder_lands_where_ini_puts_the_same_defect(self, tmp_path):
        """같은 결함을 INI 로 찍은 자리와 LIVE 파일명으로 찍은 자리가 같다(1 µm)."""
        X, Y = 150000.0, 210000.0
        ini_dir = _camtek_folder(tmp_path, [("a", X, Y)])
        (ini_pt,) = wm.build_map(resolve_batch([ini_dir / "a.jpeg"])).points
        c = camtek_ini.resolve(ini_dir / "a.jpeg")
        stem = f"TB500_RDL4 - Multi_FDV-RDL4_W1XYA1_{c.col}_{c.row}_{c.x}_{c.y}_Bump"
        live_dir = _live_folder(tmp_path, [stem])
        data = wm.build_map(resolve_batch([live_dir / f"{stem}.jpg"]))
        assert not data.unplaced and data.frame is not None
        assert data.frame.pitch_assumed is False
        assert data.frame.center_source == wm.SOURCE_OBSERVED
        (p,) = data.points
        assert (p.col, p.row) == (c.col, c.row)
        assert abs(p.x - ini_pt.x) <= 1 and abs(p.y - ini_pt.y) <= 1

    def test_no_ini_estimates_pitch_from_photos_and_says_so(self, tmp_path):
        """die 크기 파일이 없으면 TB500 상수가 아니라 **사진 좌표**로 pitch 를 추정한다.

        작은 die(4 mm) 자재 — 상수(37 mm)를 쓰면 col 30 이 웨이퍼 밖으로 나간다."""
        pytest.importorskip("PyQt6.QtWidgets")
        from aoi_verification.app import i18n
        from aoi_verification.app.ui.widgets.wafer_map_dialog import _MapPanel
        stems = ["LOT_REC_W1XYA1_30_40_3900.0_2950.0_Bump",
                 "LOT_REC_W1XYA1_2_35_100.0_200.0_Bump"]
        folder = _live_folder(tmp_path, stems, params=False)
        data = wm.build_map(resolve_batch([folder / f"{s}.jpg" for s in stems]))
        assert not data.unplaced and data.frame.pitch_assumed is True
        assert data.frame.pitch_x < 5000 and data.frame.pitch_y < 5000
        assert data.frame.pitch_x > 3900 and data.frame.pitch_y > 2950
        for p in data.points:
            assert math.hypot(p.x, p.y) <= DIA / 2
        # 칸 간격이 파일명 col/row 차이와 같다(row 는 위로 증가)
        a, b = sorted(data.points, key=lambda p: p.col)
        ka, kb = (wm.cell_of(data.frame, p.x, p.y) for p in (a, b))
        assert (kb[0] - ka[0], kb[1] - ka[1]) == (b.col - a.col, b.row - a.row)
        assert i18n.KO.WAFER_MAP_PITCH_ASSUMED in _MapPanel.legend_text(data)

    def test_ini_folder_without_pitch_is_untouched(self, tmp_path):
        """INI 항목이 있는데 검산 실패한 폴더(절대좌표)는 LIVE 기하를 쓰지 않는다."""
        # Col=Row=0 이라 상수 후보는 '의미없음', 파일 pitch 는 X 등호에서 거부된다.
        folder = _camtek_folder(tmp_path, [("a", 10000.0, 10000.0)])
        (folder / "Params_WaferInfo.ini").write_text(
            f"[Geometry]\nDieStep_X=9999.0\nDieStep_Y=9999.0\n"
            f"[Geometric]\nDiameter={DIA:.6f}\nCenter_X={CX:.6f}\nCenter_Y={CY:.6f}\n",
            encoding="utf-8")
        data = wm.build_map(resolve_batch([folder / "a.jpeg"]))
        assert data.frame.pitch_x is None and data.frame.pitch_assumed is False


class TestDefectCells:
    """결함이 든 die 칸 — 점 대신 칠할 면.  Qt 없이 순수하게 검증한다."""

    def test_cell_matches_the_grid_the_view_draws(self, tmp_path):
        """점이 든 칸의 경계는 격자선 위에 있고, 점은 그 안에 든다."""
        folder = _camtek_folder(tmp_path, [("a", 150000.0, 210000.0)])
        data = wm.build_map(resolve_batch([folder / "a.jpeg"]))
        frame = data.frame
        (p,) = data.points
        cell = wm.cell_of(frame, p.x, p.y)
        x0, y0, x1, y1 = wm.cell_bounds(frame, cell)
        assert x0 <= p.x < x1 and y0 <= p.y < y1
        xs, ys = wm.grid_lines(frame)
        for v, axis in ((x0, xs), (x1, xs), (y0, ys), (y1, ys)):
            assert min(abs(v - g) for g in axis) < 1e-6
        assert wm.defect_cells(data) == {cell}

    def test_several_defects_in_one_die_collapse_to_one_cell(self, tmp_path):
        """한 칸에 결함이 여럿이어도 칸 하나다 — 칠하는 색은 하나뿐이라 상태를 안 담는다."""
        # 같은 die 안의 두 점(칸은 같고 die 내부 위치만 다름) + 다른 die 의 점 하나.
        folder = _camtek_folder(tmp_path, [("one", 150000.0, 210000.0),
                                           ("two", 151000.0, 211000.0),
                                           ("far", 60000.0, 120000.0)])
        paths = [folder / f"{n}.jpeg" for n in ("one", "two", "far")]
        data = wm.build_map(resolve_batch(paths))
        by_stem = {p.path.stem: p for p in data.points}
        at = lambda n: wm.cell_of(data.frame, by_stem[n].x, by_stem[n].y)  # noqa: E731
        assert at("one") == at("two") != at("far")
        assert wm.defect_cells(data) == {at("one"), at("far")}   # 세 점이 두 칸

    def test_no_pitch_means_no_cells(self, tmp_path):
        """die 기하를 모르는(절대좌표) 프레임은 칸을 못 정한다 — 화면은 점으로 남는다."""
        folder = _camtek_folder(tmp_path, [("a", 150000.0, 210000.0)])
        data = wm.build_map(resolve_batch([folder / "a.jpeg"]))
        flat = wm.WaferFrame(diameter=DIA, pitch_x=None, pitch_y=None,
                             grid_x0=0.0, grid_y0=0.0,
                             center_source=wm.SOURCE_OBSERVED, kind="camtek")
        assert wm.cell_of(flat, 0.0, 0.0) is None
        assert wm.defect_cells(wm.MapData(flat, data.points, ())) == set()


class TestKlaPlane:
    def test_roundtrip_to_index_frame(self, tmp_path):
        defects = [("k1", 1234.5, 2345.6, -2, -1), ("k2", 100.0, 40000.0, 1, 2)]
        folder = _kla_folder(tmp_path, defects)
        data = wm.build_map(resolve_batch([folder / f"{s}.jpg" for s, *_ in defects]))
        assert data.frame is not None and data.frame.kind == "kla"
        assert data.frame.center_source == wm.SOURCE_OBSERVED
        by_name = {p.path.stem: p for p in data.points}
        for stem, xr, yr, xi, yi in defects:
            p = by_name[stem]
            # DefectCoord.x/y 는 round 라 0.5 µm 오차 허용.
            assert p.x == pytest.approx(xi * PX + xr - CX, abs=0.5)
            assert p.y == pytest.approx(yi * PY + yr - CY, abs=0.5)

    def test_no_center_assumed(self, tmp_path):
        folder = _kla_folder(tmp_path, [("k1", 100.0, 200.0, 0, 0)], center=False)
        data = wm.build_map(resolve_batch([folder / "k1.jpg"]))
        assert data.frame.center_source == wm.SOURCE_ASSUMED
        (p,) = data.points
        assert math.hypot(p.x, p.y) <= DIA / 2


def test_slot_maps_merges_lot_and_marks_matches(tmp_path):
    from aoi_verification.app.models.result import FinalResult, MatchResult
    f1 = _camtek_folder(tmp_path / "s1", [("a", 150000.0, 210000.0)])
    f2 = _camtek_folder(tmp_path / "s2", [("b", 60000.0, 120000.0)])
    result = FinalResult(
        mode="single", ref_machine="1", val_machine="2",
        # S1 의 매치 상대는 목록 밖의 파일 — S2 의 b 는 어느 매치에도 없다.
        matches=[MatchResult(slot="S1", ref_path=f1 / "a.jpeg",
                             val_path=f2 / "c.jpeg", score=1.0)],
        slot_images={"S1": ([f1 / "a.jpeg"], []), "S2": ([], [f2 / "b.jpeg"])},
    )
    ref, val = wm.slot_maps(result, wm.ALL_SLOTS_KEY)
    assert [p.matched for p in ref.points] == [True]
    assert [p.matched for p in val.points] == [False]
    ref1, val1 = wm.slot_maps(result, "S1")
    assert len(ref1.points) == 1 and not val1.points


def test_resolve_batch_reports_progress(tmp_path):
    """좌표 해석은 결정형 진행을 낸다 — 마지막 보고가 (n, n) 이고 done 은 단조 증가."""
    n = 120
    folder = _camtek_folder(tmp_path, [(f"p{i}", CX, CY) for i in range(n)])
    seen: list[tuple[int, int]] = []
    resolve_batch([folder / f"p{i}.jpeg" for i in range(n)],
                  progress=lambda d, t: seen.append((d, t)))
    assert seen and seen[-1] == (n, n)
    assert all(t == n for _, t in seen)
    assert [d for d, _ in seen] == sorted(d for d, _ in seen)


def test_slot_maps_progress_sums_ref_and_val(tmp_path):
    from aoi_verification.app.models.result import FinalResult
    f1 = _camtek_folder(tmp_path / "s1", [("a", CX, CY), ("b", CX, CY)])
    f2 = _camtek_folder(tmp_path / "s2", [("c", CX, CY)])
    result = FinalResult(mode="single", ref_machine="1", val_machine="2", matches=[],
                         slot_images={"S1": ([f1 / "a.jpeg", f1 / "b.jpeg"],
                                             [f2 / "c.jpeg"])})
    seen: list[tuple[int, int]] = []
    wm.slot_maps(result, "S1", progress=lambda d, t: seen.append((d, t)))
    assert seen == [(2, 3), (3, 3)]        # 기준 2장 → 검증 1장, 총량은 합산 3


# ---------------------------------------------------------------------------
# Qt — 뷰 · 시트 · 엑셀
# ---------------------------------------------------------------------------
@pytest.fixture
def qt():
    pytest.importorskip("PyQt6.QtWidgets")
    from PyQt6.QtWidgets import QApplication
    return QApplication.instance() or QApplication([])


def _view_with_points(qt, tmp_path):
    from aoi_verification.app.ui.widgets.wafer_map_view import WaferMapView
    folder = _camtek_folder(tmp_path, [("a", CX, CY), ("b", 60000.0, 120000.0)])
    data = wm.build_map(resolve_batch([folder / "a.jpeg", folder / "b.jpeg"]))
    view = WaferMapView()
    view.resize(400, 400)
    view.set_data(data)
    return view, data


def test_view_hit_test_and_wheel_zoom(qt, tmp_path):
    from PyQt6.QtCore import QPoint, QPointF, Qt
    from PyQt6.QtGui import QWheelEvent
    view, data = _view_with_points(qt, tmp_path)
    center = QPointF(200, 200)
    hit = view.point_at(center)             # 'a' 는 정확히 웨이퍼 중심
    assert hit is not None and hit.path.stem == "a"
    assert view.point_at(QPointF(20, 20)) is None
    ev = QWheelEvent(center, QPointF(view.mapToGlobal(QPoint(200, 200))), QPoint(0, 0),
                     QPoint(0, 240), Qt.MouseButton.NoButton,
                     Qt.KeyboardModifier.NoModifier, Qt.ScrollPhase.NoScrollPhase,
                     False)
    view.wheelEvent(ev)
    assert view.zoom() > 1.0
    assert view.point_at(center).path.stem == "a"     # 커서 아래 점은 그대로
    view.grab()                                        # paintEvent 가 예외 없이 돈다
    view.reset_view()
    assert view.zoom() == 1.0


def test_fill_dies_paints_the_die_not_the_dot(qt, tmp_path):
    """die 색칠: 결함이 든 **그 칸만** 칠한다 — 빈 칸은 그대로, 점은 생략."""
    from aoi_verification.app.ui.widgets.wafer_map_view import _colors
    view, data = _view_with_points(qt, tmp_path)
    frame = data.frame
    pt = next(p for p in data.points if p.path.stem == "a")
    cell = wm.cell_of(frame, pt.x, pt.y)
    x0, y0, _x1, _y1 = wm.cell_bounds(frame, cell)
    inside = ((x0 + pt.x) / 2, (y0 + pt.y) / 2)       # 칸 안, 점에서 떨어진 자리
    empty = wm.cell_bounds(frame, (cell[0], cell[1] + 1))   # 결함 없는 이웃 칸
    assert (cell[0], cell[1] + 1) not in wm.defect_cells(data)
    empty_mid = ((empty[0] + empty[2]) / 2, (empty[1] + empty[3]) / 2)
    m = view._mapper()
    at = lambda xy: m.to_px(*xy).toPoint()            # noqa: E731
    col = _colors()

    off = view.grab().toImage()
    assert off.pixelColor(at(inside)).name() == col["wafer"].name()
    dot_color = off.pixelColor(at((pt.x, pt.y))).name()      # 점 모드에서 찍힌 그 색
    assert dot_color == col["neutral"].name()
    view.set_fill_dies(True)
    assert view.fill_dies()
    on = view.grab().toImage()
    # ★ 칸 색은 **기존 결함 점 색과 같은 색**이다(사용자 결정) — 매치됨/미매치로
    #   갈리지 않는다.  켜는 곳이 셋업 단계뿐이라 거기서는 점이 전부 이 색이다.
    assert on.pixelColor(at(inside)).name() == dot_color
    assert on.pixelColor(at(empty_mid)).name() == col["wafer"].name()
    # 칠하기 모드에서는 점이 안 보이므로 칸 아무 데나 집어도 그 결함이 잡힌다.
    assert view.point_at(m.to_px(*inside)).path.stem == "a"
    assert view.point_at(m.to_px(*empty_mid)) is None
    view.set_fill_dies(False)
    assert view.point_at(m.to_px(*inside)) is None     # 점 모드는 HIT_PX 그대로


def test_rotation_turns_map_and_hit_test_clockwise(qt, tmp_path):
    """노치 회전: 점·히트 판정이 화면 시계 방향으로 같이 돈다(평면 값은 그대로)."""
    from PyQt6.QtCore import QPointF
    view, data = _view_with_points(qt, tmp_path)
    pt = next(p for p in data.points if p.path.stem == "b")
    before = view._mapper().to_px(pt.x, pt.y)
    assert view.point_at(before).path.stem == "b"
    assert view.rotation() == 0

    view.set_rotation(1)
    assert view.rotation() == 1
    after = view._mapper().to_px(pt.x, pt.y)
    c = QPointF(200, 200)                       # 화면 중앙 = 웨이퍼 중심
    # 화면은 y 가 아래로 증가 — 시계 방향 90° 는 (dx, dy) → (−dy, dx).
    assert after.x() == pytest.approx(c.x() - (before.y() - c.y()), abs=0.5)
    assert after.y() == pytest.approx(c.y() + (before.x() - c.x()), abs=0.5)
    assert view.point_at(after).path.stem == "b"
    assert view.point_at(before) is None        # 옛 자리에는 없다
    assert [p.x for p in view.data().points] == [p.x for p in data.points]

    view.set_rotation(4)                        # 한 바퀴 = 제자리
    assert view.rotation() == 0
    assert view._mapper().to_px(pt.x, pt.y) == before


def test_notch_follows_the_rotation(qt, tmp_path):
    """노치 자국이 아래에서 왼쪽으로 옮겨간다 — 원 윤곽이 실제로 돈다."""
    from aoi_verification.app.ui.widgets.wafer_map_view import _colors
    view, _ = _view_with_points(qt, tmp_path)
    bg = _colors()["bg"].name()

    def bg_near(img, cx, cy) -> int:
        """(cx, cy) 둘레 9×9 px 중 바탕색(=원 바깥/노치 파임) 픽셀 수."""
        return sum(img.pixelColor(x, y).name() == bg
                   for x in range(cx - 4, cx + 5) for y in range(cy - 4, cy + 5))

    bottom, left = (200, 382), (18, 200)        # 원 반지름 184 px, 중심 (200, 200)
    img0 = view.grab().toImage()
    view.set_rotation(1)
    img1 = view.grab().toImage()
    assert bg_near(img0, *bottom) > bg_near(img1, *bottom)      # 아래 파임이 메워지고
    assert bg_near(img1, *left) > bg_near(img0, *left)          # 왼쪽이 파인다


def test_rotate_xy_walks_the_notch_around(qt):
    """``rotate_xy`` 의 방향 규약 — i18n 의 노치 방향 이름 순서와 같아야 한다."""
    from aoi_verification.app import i18n
    from aoi_verification.app.ui.widgets.wafer_map_view import ROTATIONS, rotate_xy
    assert len(i18n.KO.WAFER_MAP_NOTCH_DIRS) == ROTATIONS
    # 노치는 평면 (0, −1) — 아래 → 왼쪽 → 위 → 오른쪽.
    assert [rotate_xy(0.0, -1.0, k) for k in range(4)] == [
        (0.0, -1.0), (-1.0, 0.0), (0.0, 1.0), (1.0, 0.0)]
    assert rotate_xy(*rotate_xy(3.0, 5.0, 1), -1) == (3.0, 5.0)     # 역변환


def test_render_png_draws_wafer(qt, tmp_path):
    from aoi_verification.app.ui.widgets.wafer_map_view import render_map_image
    _, data = _view_with_points(qt, tmp_path)
    img = render_map_image(data, size=240)
    assert not img.isNull() and img.width() == 240
    c = img.pixelColor(120, 120)
    bg = img.pixelColor(2, 2)
    assert c.name() != bg.name()          # 중심(점)은 바탕과 다른 색


def _wait_build(qt, dlg, timeout_ms: int = 10000) -> None:
    """워커가 끝나고 done 시그널이 처리될 때까지."""
    from PyQt6.QtCore import QDeadlineTimer
    dl = QDeadlineTimer(timeout_ms)
    while dlg.is_building() and not dl.hasExpired():
        qt.processEvents()
    qt.processEvents()
    qt.processEvents()
    assert not dlg.is_building(), "맵 워커가 끝나지 않았다"


def test_dialog_result_mode_shows_two_maps(qt, tmp_path):
    from aoi_verification.app.models.result import FinalResult, MatchResult
    from aoi_verification.app.ui.widgets.wafer_map_dialog import WaferMapDialog
    f1 = _camtek_folder(tmp_path / "s1", [("a", CX, CY)])
    f2 = _kla_folder(tmp_path / "s1", [("k", 100.0, 200.0, 0, 0)])
    result = FinalResult(
        mode="single", ref_machine="1호기", val_machine="KLA",
        matches=[MatchResult(slot="S1", ref_path=f1 / "a.jpeg",
                             val_path=f2 / "k.jpg", score=1.0)],
        slot_images={"S1": ([f1 / "a.jpeg"], [f2 / "k.jpg"])},
    )
    dlg = WaferMapDialog(result=result)
    try:
        assert dlg.is_building()                      # 워커가 돈다 — 로딩이 뜬 상태
        _wait_build(qt, dlg)
        assert dlg.slot_combo.count() == 2
        assert not dlg.empty.isVisibleTo(dlg)
        assert dlg.left.isVisibleTo(dlg) and dlg.right.isVisibleTo(dlg)
        assert dlg.left.view.data().points[0].matched is True
        assert dlg.right.view.data().frame.kind == "kla"
        assert "1호기" in dlg.left.title.text()
        dlg.slot_combo.setCurrentIndex(1)
        _wait_build(qt, dlg)
        assert dlg.current_slot() == "S1"
    finally:
        dlg.deleteLater()


def test_dialog_setup_mode_slot_folder(qt, tmp_path):
    from aoi_verification.app import i18n
    from aoi_verification.app.ui.widgets.wafer_map_dialog import WaferMapDialog
    dlg = WaferMapDialog()
    try:
        assert dlg.empty.isVisibleTo(dlg) and not dlg.left.isVisibleTo(dlg)
        empty = tmp_path / "empty"
        empty.mkdir()
        dlg.show_folder(empty)
        # 폴더를 고른 **그 순간** 오버레이가 무슨 일인지 말한다 — 워커 첫 보고 전에.
        assert not dlg._loading.isHidden()          # 시트 자체는 아직 show 전이라 isVisibleTo
        assert dlg._loading._label.text() == i18n.KO.WAFER_MAP_LOADING_SCAN
        _wait_build(qt, dlg)
        assert dlg.empty.text() == i18n.KO.WAFER_MAP_NO_IMAGES
        folder = _camtek_folder(tmp_path, [("a", CX, CY)])
        dlg.show_folder(folder)
        _wait_build(qt, dlg)
        assert dlg.left.isVisibleTo(dlg) and not dlg.right.isVisibleTo(dlg)
        assert not dlg.slots_btn.isVisibleTo(dlg)      # 슬롯 폴더 — 슬롯 선택 없음
        assert i18n.KO.WAFER_MAP_LEGEND_DEFECT in dlg.left.legend.text()
    finally:
        dlg.deleteLater()


def test_dialog_single_image_shows_only_that_defect(qt, tmp_path):
    """사진 1장 — 고른 결함만 찍고, 원·격자는 그 폴더 기하 그대로."""
    from aoi_verification.app import i18n
    from aoi_verification.app.ui.widgets.wafer_map_dialog import WaferMapDialog
    folder = _camtek_folder(tmp_path, [("a", CX, CY), ("b", 60000.0, 120000.0)])
    dlg = WaferMapDialog()
    try:
        dlg.show_image(folder / "b.jpeg")
        _wait_build(qt, dlg)
        data = dlg.left.view.data()
        assert [p.path.stem for p in data.points] == ["b"]
        assert data.frame is not None and data.frame.pitch_x == pytest.approx(PX)
        assert dlg.single_image() == folder / "b.jpeg"
        assert dlg.left.title.text() == i18n.KO.WAFER_MAP_ONE_IMAGE_FMT.format(
            name="b.jpeg")
        assert not dlg.slots_btn.isVisibleTo(dlg)

        stray = folder / "no_such_entry.jpeg"       # 좌표를 못 읽는 사진
        stray.write_bytes(b"")
        dlg.show_image(stray)
        _wait_build(qt, dlg)
        assert dlg.empty.isVisibleTo(dlg)
        assert dlg.empty.text() == i18n.KO.WAFER_MAP_NO_COORD

        dlg.show_folder(folder)                     # 폴더로 돌아오면 1장 상태가 풀린다
        _wait_build(qt, dlg)
        assert dlg.single_image() is None
        assert len(dlg.left.view.data().points) == 2
        assert dlg.left.title.text() == folder.name
    finally:
        dlg.deleteLater()


def test_fill_toggle_is_setup_only_and_notch_turns_both_maps(qt, tmp_path):
    """die 색칠은 **셋업 단계에만** 단다(사용자 결정) — 결과 단계는 점 색이 곧
    매치됨/미매치라 한 색으로 칠하면 그 정보가 사라진다.  노치 회전은 두 단계 모두
    이고, 결과 단계에서는 기준·검증 두 맵에 함께 걸린다."""
    from aoi_verification.app import i18n
    from aoi_verification.app.models.result import FinalResult
    from aoi_verification.app.ui.widgets.wafer_map_dialog import WaferMapDialog
    f1 = _camtek_folder(tmp_path / "s1", [("a", CX, CY)])
    f2 = _camtek_folder(tmp_path / "s2", [("b", 60000.0, 120000.0)])
    result = FinalResult(mode="single", ref_machine="1", val_machine="2", matches=[],
                         slot_images={"S1": ([f1 / "a.jpeg"], [f2 / "b.jpeg"])})
    dlg = WaferMapDialog(result=result)
    try:
        _wait_build(qt, dlg)
        views = (dlg.left.view, dlg.right.view)
        assert dlg.fill_btn is None                 # 결과 단계 — 칠하기 없음
        assert not any(v.fill_dies() for v in views)
        assert dlg.notch_btn.text() == i18n.KO.WAFER_MAP_NOTCH_FMT.format(dir="아래")
        dlg.notch_btn.click()
        assert [v.rotation() for v in views] == [1, 1]
        assert dlg.notch_btn.text() == i18n.KO.WAFER_MAP_NOTCH_FMT.format(dir="왼쪽")
        dlg.slot_combo.setCurrentIndex(1)           # 새 맵이 들어와도 보기 상태 유지
        _wait_build(qt, dlg)
        assert [v.rotation() for v in views] == [1, 1]
        for _ in range(3):                          # 네 번 누르면 제자리
            dlg.notch_btn.click()
        assert [v.rotation() for v in views] == [0, 0]
        assert dlg.notch_btn.text() == i18n.KO.WAFER_MAP_NOTCH_FMT.format(dir="아래")
    finally:
        dlg.deleteLater()

    setup = WaferMapDialog()                        # 셋업 단계 — 칠하기가 있다
    try:
        setup.show_folder(f1)
        _wait_build(qt, setup)
        assert setup.fill_btn is not None
        assert setup.fill_btn.isCheckable()
        assert not setup.left.view.fill_dies()
        setup.fill_btn.setChecked(True)
        assert setup.left.view.fill_dies()
        setup.show_folder(f2)                       # 새 맵이 들어와도 유지된다
        _wait_build(qt, setup)
        assert setup.left.view.fill_dies()
        setup.notch_btn.click()
        assert setup.left.view.rotation() == 1
    finally:
        setup.deleteLater()


def test_fullscreen_hides_the_chrome_and_grows_the_map(qt, tmp_path):
    """전체화면: 주변 표시가 빠져 맵 **몫이 커지고**, 앱 창도 같이 전체화면이 된다.

    시트와 같은 모양으로(창 안 자식 위젯) 띄워서 잰다 — 실제 경로가 그렇다.  높이 자체
    대신 **창에서 맵이 차지하는 비율**을 비교한다: 창이 전체화면으로 커지는 것과 주변
    표시가 빠지는 것을 섞지 않고, 뒤엣것만 본다."""
    from PyQt6.QtCore import QEvent, Qt
    from PyQt6.QtGui import QKeyEvent
    from PyQt6.QtWidgets import QVBoxLayout, QWidget
    from aoi_verification.app.ui.widgets.wafer_map_dialog import WaferMapDialog
    folder = _camtek_folder(tmp_path, [("a", CX, CY)])
    win = QWidget()
    win.resize(900, 700)
    lay = QVBoxLayout(win)
    lay.setContentsMargins(0, 0, 0, 0)
    dlg = WaferMapDialog(win)
    dlg.setWindowFlags(Qt.WindowType.Widget)         # 시트처럼 창 **안**으로
    lay.addWidget(dlg)
    win.show()
    qt.processEvents()
    try:
        dlg.show_folder(folder)
        _wait_build(qt, dlg)
        qt.processEvents()
        share = lambda: dlg.left.view.height() / max(1, dlg.height())   # noqa: E731
        before = share()
        assert 0 < before < 1
        assert not dlg.is_fullscreen() and not win.isFullScreen()
        assert dlg.exit_btn.isHidden()

        dlg.full_btn.click()
        qt.processEvents()
        assert dlg.is_fullscreen()
        assert win.isFullScreen(), "앱 창이 전체화면으로 바뀌지 않았다"
        # 주변 표시가 전부 빠지고 맵만 남는다.
        assert not dlg.sub.isVisibleTo(dlg)
        assert not dlg._top_host.isVisibleTo(dlg)
        assert not dlg.left.title.isVisibleTo(dlg)
        assert not dlg.left.legend.isVisibleTo(dlg)
        assert dlg.layout().contentsMargins().top() == 0
        assert dlg.exit_btn.isVisibleTo(dlg)         # 유일한 나가기 버튼이 뜬다
        assert dlg.exit_btn.x() + dlg.exit_btn.width() <= dlg.width()
        full_share = share()
        assert full_share > before, "맵 몫이 커지지 않았다"

        dlg.keyPressEvent(QKeyEvent(QEvent.Type.KeyPress, Qt.Key.Key_Escape,
                                    Qt.KeyboardModifier.NoModifier))
        qt.processEvents()
        assert not dlg.is_fullscreen()
        assert not win.isFullScreen(), "창이 되돌아오지 않았다"
        assert dlg.sub.isVisibleTo(dlg) and dlg._top_host.isVisibleTo(dlg)
        assert dlg.left.title.isVisibleTo(dlg)
        assert dlg.exit_btn.isHidden()
        # 창 크기가 돌아오는 폭은 플랫폼마다 달라 절대값은 안 본다 — 주변 표시가
        # 되살아나 맵 몫이 **다시 줄었다**는 것만 본다.
        assert share() < full_share, "주변 표시가 되살아나지 않았다"

        dlg.full_btn.click()                         # 떠 있는 버튼으로도 나온다
        assert dlg.is_fullscreen()
        dlg.exit_btn.click()
        qt.processEvents()
        assert not dlg.is_fullscreen() and not win.isFullScreen()
    finally:
        win.hide()
        qt.processEvents()
        win.deleteLater()
        qt.processEvents()


def test_fullscreen_leaves_an_unshown_window_alone(qt, tmp_path):
    """안 보이는 창은 건드리지 않는다 — 뜰 생각이 없던 창을 띄우면 안 된다.

    닫을 때는 전체화면을 스스로 푼다(내가 키운 창이 전체화면으로 남으면 안 된다)."""
    from aoi_verification.app.ui.widgets.wafer_map_dialog import WaferMapDialog
    dlg = WaferMapDialog()
    try:
        dlg.set_fullscreen(True)
        assert dlg.is_fullscreen()
        assert not dlg.isVisible() and not dlg.isFullScreen()
        assert dlg._win_changed is False           # 창 상태는 손대지 않았다
        dlg.close()
        assert not dlg.is_fullscreen()
    finally:
        dlg.deleteLater()


def test_excel_map_ignores_view_options(qt, tmp_path):
    """엑셀에 들어가는 맵은 화면 보기 옵션을 따라가지 않는다(사용자 결정) — 결과
    파일이 '언제 내보냈는지' 에 따라 달라지면 안 된다."""
    from aoi_verification.app.ui.widgets.wafer_map_view import render_map_image
    view, data = _view_with_points(qt, tmp_path)
    before = render_map_image(data, size=240)
    view.set_rotation(2)
    view.set_fill_dies(True)
    assert render_map_image(data, size=240) == before


def test_dialog_overlay_gets_determinate_coord_progress(qt, tmp_path, monkeypatch):
    """좌표 읽기 단계가 오버레이에 (done, total>0) 로 도착한다 — busy 로만 머물면
    수천 장 파싱 동안 '진척도가 안 나타난다'(실제 보고)."""
    from aoi_verification.app.ui.widgets.wafer_map_dialog import WaferMapDialog
    folder = _camtek_folder(tmp_path, [(f"p{i}", CX, CY) for i in range(60)])
    dlg = WaferMapDialog()
    calls: list[tuple[int, int, str]] = []
    real = dlg._loading.set_progress
    monkeypatch.setattr(dlg._loading, "set_progress",
                        lambda d, t, m="": (calls.append((d, t, m)), real(d, t, m)))
    try:
        dlg.show_folder(folder)
        _wait_build(qt, dlg)
        det = [(d, t) for d, t, _ in calls if t > 0]
        assert (60, 60) in det, calls
    finally:
        dlg.deleteLater()


def test_dialog_setup_mode_lot_folder_and_subset(qt, tmp_path):
    """LOT 폴더 → 전체 합산 먼저, 슬롯 선택으로 일부만."""
    from aoi_verification.app import i18n
    from aoi_verification.app.ui.widgets.wafer_map_dialog import (WaferMapDialog,
                                                                 classify_folder)
    lot = tmp_path / "LOT1"
    _camtek_folder(lot, [("a", CX, CY)]).rename(lot / "S1")
    _camtek_folder(lot, [("b", 60000.0, 120000.0)]).rename(lot / "S2")
    (lot / "S3").mkdir()                                  # 사진 없는 폴더는 슬롯이 아니다
    kind, slots = classify_folder(lot)
    assert kind == "lot" and set(slots) == {"S1", "S2"}   # 슬롯명 = LOT 바로 아래 폴더명
    dlg = WaferMapDialog()
    try:
        dlg.show_folder(lot)
        _wait_build(qt, dlg)
        assert dlg.slots_btn.isVisibleTo(dlg)
        assert len(dlg.left.view.data().points) == 2
        assert dlg.left.title.text() == i18n.KO.WAFER_MAP_LOT_ALL_FMT.format(
            lot="LOT1", total=2)
        dlg._selected = {"S2"}
        dlg._rebuild_folder_map()
        _wait_build(qt, dlg)
        pts = dlg.left.view.data().points
        assert [p.path.stem for p in pts] == ["b"]
        assert "1/2" in dlg.left.title.text()
    finally:
        dlg.deleteLater()


def test_build_prewarms_thumbnails(qt, isolated_cache, tmp_path):
    pytest.importorskip("PIL.Image")
    from PIL import Image
    from aoi_verification.app.ui.widgets.wafer_map_dialog import WaferMapDialog
    from aoi_verification.app.utils import image_io
    folder = _camtek_folder(tmp_path, [("a", CX, CY)])
    Image.new("RGB", (400, 300)).save(str(folder / "a.jpeg"), "JPEG")
    dlg = WaferMapDialog()
    try:
        dlg.show_folder(folder)
        _wait_build(qt, dlg)
        thumb = image_io.get_map_thumb_path(folder / "a.jpeg")   # 이미 있어야 한다
        assert thumb.exists() and thumb.stat().st_size > 0
        # 저화질 전용 — 일반 썸네일(240px)과 다른 파일, 긴 변이 MAP_THUMB_PX 이하.
        from aoi_verification.app.config import Sizing
        assert thumb != image_io.get_thumb_path(folder / "a.jpeg")
        assert max(Image.open(thumb).size) <= Sizing.MAP_THUMB_PX
    finally:
        dlg.deleteLater()


def test_prewarm_skipped_above_cap(qt, isolated_cache, tmp_path, monkeypatch):
    """사진이 PREWARM_MAX 장을 넘으면 썸네일 선로딩을 건너뛴다(사용자 결정: 500)."""
    pytest.importorskip("PIL.Image")
    from PIL import Image
    from aoi_verification.app.ui.widgets import wafer_map_dialog as wmd
    from aoi_verification.app.utils import image_io
    assert wmd.PREWARM_MAX == 500
    monkeypatch.setattr(wmd, "PREWARM_MAX", 1)
    folder = _camtek_folder(tmp_path, [("a", CX, CY), ("b", 60000.0, 120000.0)])
    for n in ("a", "b"):
        Image.new("RGB", (40, 30)).save(str(folder / f"{n}.jpeg"), "JPEG")
    calls = []
    monkeypatch.setattr(image_io, "get_map_thumb_path",
                        lambda p, **k: calls.append(p) or p)
    dlg = wmd.WaferMapDialog()
    try:
        dlg.show_folder(folder)
        _wait_build(qt, dlg)
        assert len(dlg.left.view.data().points) == 2
        assert calls == []
    finally:
        dlg.deleteLater()


def test_view_double_click_activates_point(qt, tmp_path):
    from PyQt6.QtCore import QEvent, QPointF, Qt
    from PyQt6.QtGui import QMouseEvent
    view, _ = _view_with_points(qt, tmp_path)
    got = []
    view.point_activated.connect(got.append)
    center = QPointF(200, 200)
    mk = lambda kind: QMouseEvent(kind, center, Qt.MouseButton.LeftButton,   # noqa: E731
                                  Qt.MouseButton.LeftButton,
                                  Qt.KeyboardModifier.NoModifier)
    view.mousePressEvent(mk(QEvent.Type.MouseButtonPress))
    view.mouseReleaseEvent(mk(QEvent.Type.MouseButtonRelease))
    assert got == []                                         # 단일 클릭은 아무것도 안 연다
    view.mouseDoubleClickEvent(mk(QEvent.Type.MouseButtonDblClick))
    assert [p.stem for p in got] == ["a"]
    view.mouseDoubleClickEvent(QMouseEvent(
        QEvent.Type.MouseButtonDblClick, QPointF(20, 20), Qt.MouseButton.LeftButton,
        Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier))
    assert len(got) == 1                                     # 빈 곳은 원래 크기 복귀만


def test_setup_and_result_pages_have_buttons(qt):
    from aoi_verification.app import i18n
    from aoi_verification.app.ui.pages import result_page as rp, setup_page as sp
    page = sp.SetupPage()
    try:
        bar = page._action_bar
        assert bar.itemAt(2).widget() is page.wafer_map_btn      # 사진 정보 보기 옆
        assert bar.itemAt(bar.count() - 1).widget() is page.start_btn
        assert page.wafer_map_btn.text() == i18n.KO.WAFER_MAP_BUTTON
    finally:
        page.deleteLater()
    res = rp.ResultPage()
    try:
        assert res.wafer_map_btn.text() == i18n.KO.WAFER_MAP_BUTTON
    finally:
        res.deleteLater()


def test_export_writes_wafer_map_sheet(qt, isolated_cache, tmp_path):
    pytest.importorskip("openpyxl")
    pytest.importorskip("PIL.Image")
    from openpyxl import load_workbook
    from PIL import Image

    from aoi_verification.app import i18n
    from aoi_verification.app.models.result import FinalResult, MatchResult
    from aoi_verification.app.workers.exporter import ExcelExporter

    f1 = _camtek_folder(tmp_path / "s1", [("a", CX, CY)])
    f2 = _camtek_folder(tmp_path / "s2", [("b", 60000.0, 120000.0)])
    for p in (f1 / "a.jpeg", f2 / "b.jpeg"):
        Image.new("RGB", (40, 30)).save(str(p), "JPEG")
    result = FinalResult(
        mode="single", ref_machine="1호기", val_machine="2호기",
        matches=[MatchResult(slot="S1", ref_path=f1 / "a.jpeg",
                             val_path=f2 / "b.jpeg", score=1.0)],
        slot_images={"S1": ([f1 / "a.jpeg"], [f2 / "b.jpeg"]),
                     "S2": ([f1 / "a.jpeg"], [])},
    )
    dst = tmp_path / "out.xlsx"
    from aoi_verification.app.ui.widgets.wafer_map_view import render_map_png
    exp = ExcelExporter(result, dst_path=dst, template_path=tmp_path / "none.xlsx",
                        map_renderer=render_map_png)
    exp.run()
    wb = load_workbook(str(dst))
    ws = wb[i18n.KO.WAFER_MAP_SHEET]
    assert [ws.cell(row=r, column=1).value for r in (2, 3, 4)] == [
        i18n.KO.WAFER_MAP_SHEET_ALL, "S1", "S2"]
    # 전체(기준·검증) + S1(기준·검증) + S2(기준만) = 5 장.
    assert len(ws._images) == 5


def test_wafer_map_sheet_names_the_machines_and_slot_numbers(
        qt, isolated_cache, tmp_path):
    """Wafer Map 시트도 '어느 호기의 맵인지' 와 '몇 번 슬롯인지' 를 적는다.

    요약 시트는 머리 2행(그룹 / AOI-N)과 B열 `(#6)` 으로 둘 다 밝히는데, 이 시트만
    `기준`·`검증`·슬롯명뿐이라 호기와 카세트 번호가 빠져 있었다(사용자 지적).
    """
    pytest.importorskip("openpyxl")
    pytest.importorskip("PIL.Image")
    from openpyxl import load_workbook
    from PIL import Image

    from aoi_verification.app import i18n
    from aoi_verification.app.models.result import FinalResult, MatchResult
    from aoi_verification.app.workers.exporter import ExcelExporter

    f1 = _camtek_folder(tmp_path / "s1", [("a", CX, CY)])
    f2 = _camtek_folder(tmp_path / "s2", [("b", 60000.0, 120000.0)])
    for p in (f1 / "a.jpeg", f2 / "b.jpeg"):
        Image.new("RGB", (40, 30)).save(str(p), "JPEG")
    result = FinalResult(
        mode="cross", ref_machine="17", val_machine="23",
        matches=[MatchResult(slot="A1033ABQEWG3", ref_path=f1 / "a.jpeg",
                             val_path=f2 / "b.jpeg", score=1.0)],
        slot_images={"A1033ABQEWG3": ([f1 / "a.jpeg"], [f2 / "b.jpeg"])},
        slot_numbers={"A1033ABQEWG3": "6"},
    )
    dst = tmp_path / "out.xlsx"
    from aoi_verification.app.ui.widgets.wafer_map_view import render_map_png
    ExcelExporter(result, dst_path=dst, template_path=tmp_path / "none.xlsx",
                  map_renderer=render_map_png).run()
    ws = load_workbook(str(dst))[i18n.KO.WAFER_MAP_SHEET]

    # 머리칸 — 역할 + 호기(요약 시트의 'AOI-N' 과 같은 라벨).
    assert ws["B1"].value == f"{i18n.KO.WAFER_MAP_SHEET_COL_REF}\nAOI-17"
    assert ws["C1"].value == f"{i18n.KO.WAFER_MAP_SHEET_COL_VAL}\nAOI-23"
    # 줄바꿈이 보이려면 wrap_text 가 있어야 한다.
    assert ws["B1"].alignment.wrap_text is True
    # 슬롯 칸 — 요약 시트 B열과 같은 표기(slot명 아래 `(#6)`).
    assert ws["A2"].value == "A1033ABQEWG3\n(#6)"
    assert ws["A2"].alignment.wrap_text is True


def test_wafer_map_sheet_header_without_machine_input(qt, isolated_cache, tmp_path):
    """호기 입력이 비어 있으면 역할만 적는다 — 빈 줄을 남기지 않는다."""
    pytest.importorskip("openpyxl")
    pytest.importorskip("PIL.Image")
    from openpyxl import load_workbook
    from PIL import Image

    from aoi_verification.app import i18n
    from aoi_verification.app.models.result import FinalResult
    from aoi_verification.app.workers.exporter import ExcelExporter

    f1 = _camtek_folder(tmp_path / "s1", [("a", CX, CY)])
    Image.new("RGB", (40, 30)).save(str(f1 / "a.jpeg"), "JPEG")
    result = FinalResult(mode="single", ref_machine="", val_machine="",
                         slot_images={"S1": ([f1 / "a.jpeg"], [])})
    dst = tmp_path / "out.xlsx"
    from aoi_verification.app.ui.widgets.wafer_map_view import render_map_png
    ExcelExporter(result, dst_path=dst, template_path=tmp_path / "none.xlsx",
                  map_renderer=render_map_png).run()
    ws = load_workbook(str(dst))[i18n.KO.WAFER_MAP_SHEET]
    assert ws["B1"].value == i18n.KO.WAFER_MAP_SHEET_COL_REF
    assert ws["C1"].value == i18n.KO.WAFER_MAP_SHEET_COL_VAL
