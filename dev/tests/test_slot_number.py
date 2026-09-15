"""결과 엑셀 slot# 칸 — slot명 아래 카세트 슬롯 번호 `(#6)` (사용자 요청).

    "WaferInfo.ini 에 ActiveSlot=6 가 있음.  이 숫자를 이용해서 결과 엑셀의
     slot# 에서 예를들어 A1033ABQEWG3 이렇게 써있으면
     A1033ABQEWG3 / (#6) 이렇게 나오도록 해줘.
     그리고 결과 엑셀에서 slot# 지금 약간 잘려서 보이는데 전체 다 보이도록
     열 너비 조정해줘"

⚠ `WaferInfo.ini` 는 `Params_WaferInfo.ini` 와 **다른 파일**이다(`models/lot_info.py`
  의 경고 참고) — 헷갈리면 조용히 아무것도 못 읽는다.
"""

from __future__ import annotations

import os
import shutil
from pathlib import Path

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from aoi_verification.app.models.lot_info import read_active_slot   # noqa: E402

_REPO = Path(__file__).resolve().parents[2]

# 실물 slot 명(WaferID)은 12자다 — 열 폭 회귀 가드가 이 길이를 기준으로 본다.
_WAFER_ID = "A1033ABQEWG3"


def _wafer_info(folder: Path, line: str) -> None:
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "WaferInfo.ini").write_text(
        f"[AutoCycleInfo]\nAutoCycleScan=1\nInputLot=TBD-PIDS3\n{line}\n"
        f"InputWaferID={_WAFER_ID}\n", encoding="utf-8")


# ---------------------------------------------------------------------------
# 1) ActiveSlot 읽기
# ---------------------------------------------------------------------------
def test_reads_active_slot(tmp_path):
    _wafer_info(tmp_path / "s1", "ActiveSlot=6")
    assert read_active_slot(tmp_path / "s1") == "6"


def test_reads_the_real_sample_in_the_repo(tmp_path):
    """저장소의 실물 샘플(`docs/WaferInfo.ini`, ActiveSlot=5)로 확인한다."""
    sample = _REPO / "docs" / "WaferInfo.ini"
    assert sample.is_file(), "실물 샘플이 사라졌다"
    slot = tmp_path / "s_real"
    slot.mkdir()
    shutil.copy(sample, slot / "WaferInfo.ini")
    assert read_active_slot(slot) == "5"


def test_parent_folder_is_not_read(tmp_path):
    """부모의 INI 는 **다른 웨이퍼**의 번호다 — 그 슬롯에 붙이면 안 된다.

    자재·Layer(`read_lot_info`)는 로트 공통이라 부모까지 보지만, ActiveSlot 은
    웨이퍼 한 장의 값이다.  못 읽으면 번호 없이 slot명만 찍히는 게 맞다.
    """
    _wafer_info(tmp_path, "ActiveSlot=6")
    (tmp_path / "s_no_ini").mkdir()
    assert read_active_slot(tmp_path / "s_no_ini") is None


def test_missing_or_broken_is_none(tmp_path):
    assert read_active_slot(tmp_path / "nope") is None
    _wafer_info(tmp_path / "s_empty", "ActiveSlot=")
    assert read_active_slot(tmp_path / "s_empty") is None


# ---------------------------------------------------------------------------
# 2) 엑셀 B열 표기 + 열 폭
# ---------------------------------------------------------------------------
@pytest.fixture
def exported(tmp_path):
    """slot 번호가 있는 결과를 실제 exporter 로 한 번 굽는다."""
    pytest.importorskip("PyQt6.QtWidgets")
    pytest.importorskip("openpyxl")
    Image = pytest.importorskip("PIL.Image")
    from PyQt6.QtWidgets import QApplication
    from aoi_verification.app.models.result import FinalResult, MatchResult
    from aoi_verification.app.workers.exporter import ExcelExporter

    QApplication.instance() or QApplication([])
    src = tmp_path / "src"
    src.mkdir()

    def _img(name: str) -> Path:
        p = src / name
        Image.new("RGB", (120, 120), (90, 90, 90)).save(str(p), "JPEG")
        return p

    result = FinalResult(
        mode="cross", ref_machine="17", val_machine="23",
        matches=[MatchResult(slot=_WAFER_ID, ref_path=_img("r.jpg"),
                             val_path=_img("v.jpg"), score=0.9),
                 MatchResult(slot="NO_NUMBER_SLOT", ref_path=_img("r2.jpg"),
                             val_path=_img("v2.jpg"), score=0.9)],
        slot_numbers={_WAFER_ID: "6"},
        kla_folders={},
    )
    dst = tmp_path / "out.xlsx"
    # 템플릿 없이(최소 헤더) 굽는다 — B열 표기·폭은 양식 유무와 무관해야 한다.
    ExcelExporter(result, dst_path=dst,
                  template_path=tmp_path / "no_template.xlsx").run()
    from openpyxl import load_workbook
    return load_workbook(str(dst), rich_text=True)["out"]


def test_slot_cell_shows_the_number_below_the_name(exported):
    """`A1033ABQEWG3` 아래 줄에 `(#6)` — 사용자가 준 예시 그대로."""
    assert str(exported["B3"].value) == f"{_WAFER_ID}\n(#6)"
    # 줄바꿈이 보이려면 wrap_text 가 켜져 있어야 한다(없으면 한 줄로 붙어 보인다).
    assert exported["B3"].alignment.wrap_text is True


def test_slot_without_number_is_unchanged(exported):
    """번호를 못 읽은 슬롯은 예전처럼 slot명만 — 빈 `(#)` 를 찍지 않는다."""
    assert exported["B4"].value == "NO_NUMBER_SLOT"


def test_slot_column_is_wide_enough_for_a_wafer_id(exported):
    """slot# 열이 WaferID 12자를 다 보여준다(폭 9.5 에서 잘려 보였다)."""
    width = exported.column_dimensions["B"].width
    assert width is not None and width >= len(_WAFER_ID) + 1, (
        f"B열 폭 {width} — WaferID {len(_WAFER_ID)}자가 잘린다")


def test_kla_folder_line_still_comes_after_the_number(tmp_path):
    """KLA 회색 줄은 번호 **아래**에 남는다 — 둘 다 있을 때 순서가 뒤집히지 않는다."""
    pytest.importorskip("PyQt6.QtWidgets")
    pytest.importorskip("openpyxl")
    Image = pytest.importorskip("PIL.Image")
    from PyQt6.QtWidgets import QApplication
    from aoi_verification.app.models.result import FinalResult, MatchResult
    from aoi_verification.app.workers.exporter import ExcelExporter

    QApplication.instance() or QApplication([])
    src = tmp_path / "src2"
    src.mkdir()
    paths = []
    for n in ("r.jpg", "v.jpg"):
        p = src / n
        Image.new("RGB", (120, 120), (90, 90, 90)).save(str(p), "JPEG")
        paths.append(p)
    result = FinalResult(
        mode="cross", ref_machine="17", val_machine="23",
        matches=[MatchResult(slot=_WAFER_ID, ref_path=paths[0],
                             val_path=paths[1], score=0.9)],
        slot_numbers={_WAFER_ID: "6"},
        kla_folders={_WAFER_ID: "KLA_FOLDER_1"},
    )
    dst = tmp_path / "out2.xlsx"
    ExcelExporter(result, dst_path=dst,
                  template_path=tmp_path / "no_template.xlsx").run()
    from openpyxl import load_workbook
    cell = load_workbook(str(dst), rich_text=True)["out2"]["B3"]
    text = "".join(getattr(b, "text", str(b)) for b in cell.value) \
        if not isinstance(cell.value, str) else cell.value
    assert text == f"{_WAFER_ID}\n(#6)\nKLA_FOLDER_1"


# ---------------------------------------------------------------------------
# 3) 슬롯 → 번호 수집 (기준·검증 폴더 양쪽)
# ---------------------------------------------------------------------------
def _scan_of(slots):
    from aoi_verification.app.models.slot import ScanResult
    return ScanResult(slots=slots, ref_only=[], val_only=[])


def _slot(name, ref_dir=None, val_dir=None):
    from aoi_verification.app.models.slot import Slot
    return Slot(name=name, ref_dir=ref_dir, val_dir=val_dir)


def _collect(scan):
    """`MainWindow._slot_numbers` 를 창 없이 돌린다 — `self._scan` 만 본다."""
    import types
    mw = pytest.importorskip("aoi_verification.app.ui.main_window")
    return mw.MainWindow._slot_numbers(types.SimpleNamespace(_scan=scan))


def test_collects_number_from_either_side(tmp_path):
    """기준 폴더에만 INI 가 있어도 번호를 얻는다(검증 쪽만 있어도 같다)."""
    _wafer_info(tmp_path / "ref" / _WAFER_ID, "ActiveSlot=6")
    (tmp_path / "val" / _WAFER_ID).mkdir(parents=True)
    got = _collect(_scan_of({_WAFER_ID: _slot(_WAFER_ID,
                                              tmp_path / "ref" / _WAFER_ID,
                                              tmp_path / "val" / _WAFER_ID)}))
    assert got == {_WAFER_ID: "6"}


def test_both_sides_agreeing_is_written_once(tmp_path):
    _wafer_info(tmp_path / "ref" / _WAFER_ID, "ActiveSlot=6")
    _wafer_info(tmp_path / "val" / _WAFER_ID, "ActiveSlot=6")
    got = _collect(_scan_of({_WAFER_ID: _slot(_WAFER_ID,
                                              tmp_path / "ref" / _WAFER_ID,
                                              tmp_path / "val" / _WAFER_ID)}))
    assert got == {_WAFER_ID: "6"}


def test_both_sides_disagreeing_are_written_together(tmp_path):
    """두 장비에서 카세트 위치가 다를 수 있다 — 한쪽만 남기면 어느 쪽인지 알 수 없다
    (KLA 폴더명 표기와 같은 관습)."""
    _wafer_info(tmp_path / "ref" / _WAFER_ID, "ActiveSlot=6")
    _wafer_info(tmp_path / "val" / _WAFER_ID, "ActiveSlot=7")
    got = _collect(_scan_of({_WAFER_ID: _slot(_WAFER_ID,
                                              tmp_path / "ref" / _WAFER_ID,
                                              tmp_path / "val" / _WAFER_ID)}))
    assert got == {_WAFER_ID: "6 / 7"}


def test_slot_without_ini_is_absent(tmp_path):
    """번호가 없으면 키 자체가 없다 — 엑셀이 빈 `(#)` 를 찍지 않게."""
    (tmp_path / "ref" / "S1").mkdir(parents=True)
    got = _collect(_scan_of({"S1": _slot("S1", tmp_path / "ref" / "S1", None)}))
    assert got == {}
