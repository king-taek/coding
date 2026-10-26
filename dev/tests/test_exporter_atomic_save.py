"""결과 엑셀 저장의 안전장치 — 원자적 저장, 사라진 사진, 시트명, rels 청소.

- 저장은 같은 폴더 임시 파일에 완성한 뒤 `os.replace` 로 바꾼다 → 실패해도 이전 결과가 산다.
- 사진이 저장 전에 사라져도 그 한 장만 글자로 대체되고 저장은 계속된다.
- 시트명은 대소문자 무관으로 겹치지 않고, 앞뒤 작은따옴표가 없고, 31자 이하다.
- rels 청소는 Type="http://…" 가 먼저 오는 실제 형태에서도 customXml 줄을 지운다.
"""
from __future__ import annotations

import os
import zipfile
from pathlib import Path

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
pytest.importorskip("PyQt6.QtWidgets")
pytest.importorskip("openpyxl")
pytest.importorskip("PIL.Image")

from openpyxl import Workbook, load_workbook                    # noqa: E402
from PIL import Image                                           # noqa: E402

from aoi_verification.app.models.result import FinalResult, MatchResult  # noqa: E402
from aoi_verification.app.workers import exporter as ex        # noqa: E402


def _result(tmp_path):
    a, b = tmp_path / "a.jpg", tmp_path / "b.jpg"
    Image.new("RGB", (120, 90), (10, 80, 40)).save(a)
    Image.new("RGB", (120, 90), (80, 10, 40)).save(b)
    return FinalResult(mode="single", ref_machine="1", val_machine="2",
                       matches=[MatchResult(slot="S1", ref_path=a, val_path=b, score=1)],
                       unmatched_refs=[])


def _run(result, dst, **kw):
    e = ex.ExcelExporter(result, dst_path=dst, template_path=dst.parent / "none.xlsx", **kw)
    errors: list[str] = []
    e.signals.failed.connect(errors.append)
    e.run()
    return errors


def _leftovers(folder: Path) -> list[str]:
    return sorted(p.name for p in folder.iterdir()
                  if p.name.endswith((".part", ".tmp")))


def test_failed_save_keeps_previous_result_and_leaves_no_temp(
        qapp, isolated_cache, tmp_path, monkeypatch):
    out = tmp_path / "out" / "result.xlsx"
    out.parent.mkdir()
    out.write_bytes(b"PREVIOUS RESULT")

    from openpyxl.workbook.workbook import Workbook as WB
    real_save = WB.save

    def boom(self, filename):
        real_save(self, filename)               # 임시 파일은 실제로 쓰인다
        raise OSError("disk full")              # 그 뒤 실패
    monkeypatch.setattr(WB, "save", boom)

    errors = _run(_result(tmp_path), out)
    assert errors and "disk full" in errors[0]
    assert out.read_bytes() == b"PREVIOUS RESULT"
    assert _leftovers(out.parent) == []


def test_cancel_during_save_keeps_previous_result(qapp, isolated_cache, tmp_path,
                                                  monkeypatch):
    out = tmp_path / "result.xlsx"
    out.write_bytes(b"PREVIOUS RESULT")
    e = ex.ExcelExporter(_result(tmp_path), dst_path=out,
                         template_path=tmp_path / "none.xlsx")
    real = ex._strip_corporate_metadata

    def stop_then_strip(p):
        e.stop()                                # 저장이 끝난 직후 앱이 닫힌다
        real(p)
    monkeypatch.setattr(ex, "_strip_corporate_metadata", stop_then_strip)
    e.run()
    assert out.read_bytes() == b"PREVIOUS RESULT"
    assert _leftovers(tmp_path) == []


def test_successful_save_replaces_file_and_leaves_no_temp(qapp, isolated_cache, tmp_path):
    out = tmp_path / "result.xlsx"
    out.write_bytes(b"old")
    assert _run(_result(tmp_path), out) == []
    assert load_workbook(str(out)).sheetnames
    assert _leftovers(tmp_path) == []


def test_replace_permission_error_keeps_locked_file_wording(
        qapp, isolated_cache, tmp_path, monkeypatch):
    """엑셀에 열린 결과는 os.replace 가 WinError 5 로 실패한다 — UI 가 '열려 있음'
    안내를 고르는 단서('permission'/'errno 13')가 문구에 남아야 한다."""
    out = tmp_path / "result.xlsx"
    out.write_bytes(b"OPEN IN EXCEL")

    def deny(src, dst):
        raise PermissionError(13, "Access is denied", str(dst))
    monkeypatch.setattr(ex.os, "replace", deny)
    errors = _run(_result(tmp_path), out)
    assert errors and "permission" in errors[0].lower() and "errno 13" in errors[0].lower()
    assert out.read_bytes() == b"OPEN IN EXCEL"
    assert _leftovers(tmp_path) == []


def test_vanished_photo_does_not_abort_save(qapp, isolated_cache, tmp_path, monkeypatch):
    """임베드할 사진을 열어 둔 뒤(경로가 아니라 바이트) 사라져도 저장은 끝까지 간다."""
    result = _result(tmp_path)
    out = tmp_path / "result.xlsx"
    # 사진 한 장은 배치 시점에 이미 없다 → 그 칸만 글자로 대체
    result.matches[0].val_path.unlink()
    assert _run(result, out) == []
    ws = load_workbook(str(out)).worksheets[0]
    assert len(ws._images) == 1                 # 남은 한 장은 들어갔다

    # 배치는 성공했는데 저장 직전 파일이 사라지는 경우 — 경로 참조면 wb.save 에서 터졌다
    result2 = _result(tmp_path)
    out2 = tmp_path / "result2.xlsx"
    real = ex.ExcelExporter._place_image

    def place_then_delete(self, ws, src, *a, **k):
        ok = real(self, ws, src, *a, **k)
        for p in (self._embed_image_path(Path(src)),):
            p.unlink(missing_ok=True)           # 캐시 파일이 저장 전에 사라진다
        return ok
    monkeypatch.setattr(ex.ExcelExporter, "_place_image", place_then_delete)
    assert _run(result2, out2) == []
    assert len(load_workbook(str(out2)).worksheets[0]._images) == 2


# ---------------------------------------------------------------- 시트명
def test_safe_sheet_title_dedupes_case_insensitively():
    taken = {"Recipe A"}
    t = ex._safe_sheet_title("RECIPE a", taken)
    assert t.casefold() != "recipe a" and t == "RECIPE a (2)"
    taken.add(t)
    assert ex._safe_sheet_title("recipe A", taken) == "recipe A (3)"


def test_safe_sheet_title_limits_and_strips_quotes():
    long = "x" * 40
    t1 = ex._safe_sheet_title(long, set())
    assert len(t1) == 31
    t2 = ex._safe_sheet_title(long, {t1})
    assert len(t2) <= 31 and t2.endswith("(2)") and t2.casefold() != t1.casefold()
    assert ex._safe_sheet_title("'abc'", set()) == "abc"
    assert ex._safe_sheet_title("''", set()) == "Recipe"
    # 31자로 자른 끝이 작은따옴표/공백이면 그것도 정리
    t3 = ex._safe_sheet_title("a" * 30 + "'tail", set())
    assert not t3.endswith("'") and len(t3) <= 31
    # 잘라낸 뒤 접미사를 붙여도 앞뒤 따옴표가 없다
    t4 = ex._safe_sheet_title("a" * 28 + "' x", {"a" * 28 + "' x"})
    assert not t4.startswith("'") and "'" != t4[-1] and len(t4) <= 31


def test_summary_sheet_name_strips_quotes_and_ignores_case(qapp, tmp_path):
    e = ex.ExcelExporter(_result(tmp_path), dst_path=tmp_path / "'결과'.xlsx")
    assert e._summary_sheet_name() == "결과"
    e2 = ex.ExcelExporter(_result(tmp_path),
                          dst_path=tmp_path / (ex.SHEET_FULL_NAME.upper() + ".xlsx"))
    assert e2._summary_sheet_name().casefold() != ex.SHEET_FULL_NAME.casefold()


# ---------------------------------------------------------------- rels 청소
RELS = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
    'officeDocument" Target="xl/workbook.xml" Id="rId1"/>'
    '<Relationship Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
    'customXml" Target="../customXml/item1.xml" Id="rId2"/>'
    '<Relationship Target="/docMetadata/LabelInfo.xml" Id="rId3" '
    'Type="http://schemas.microsoft.com/office/2020/02/relationships/classificationlabels"/>'
    '<Relationship Type="http://schemas.openxmlformats.org/package/2006/relationships/'
    'metadata/core-properties" Target="docProps/core.xml" Id="rId4"/>'
    '</Relationships>')


def test_strip_relationship_lines_handles_type_before_target():
    out = ex._strip_relationship_lines(RELS)
    assert "customXml" not in out and "docMetadata" not in out
    assert 'Target="xl/workbook.xml"' in out and 'Target="docProps/core.xml"' in out
    assert out.count("<Relationship ") == 2


def _xlsx_with_metadata(path: Path) -> None:
    wb = Workbook()
    wb.save(path)
    with zipfile.ZipFile(path) as z:
        items = {n: z.read(n) for n in z.namelist()}
    items["customXml/item1.xml"] = b"<x/>"
    items["_rels/.rels"] = RELS.encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n, d in items.items():
            z.writestr(n, d)


def test_strip_corporate_metadata_removes_rels_and_parts(tmp_path):
    p = tmp_path / "m.xlsx"
    _xlsx_with_metadata(p)
    ex._strip_corporate_metadata(p)
    with zipfile.ZipFile(p) as z:
        assert not any(n.startswith("customXml/") for n in z.namelist())
        assert b"customXml" not in z.read("_rels/.rels")
    assert _leftovers(tmp_path) == []


def test_strip_corporate_metadata_failure_removes_clean_tmp(tmp_path, monkeypatch):
    p = tmp_path / "m.xlsx"
    _xlsx_with_metadata(p)
    real = ex._strip_relationship_lines

    def boom(text):
        raise RuntimeError("x")
    monkeypatch.setattr(ex, "_strip_relationship_lines", boom)
    with pytest.raises(RuntimeError):
        ex._strip_corporate_metadata(p)
    monkeypatch.setattr(ex, "_strip_relationship_lines", real)
    assert _leftovers(tmp_path) == []
    assert zipfile.is_zipfile(p)


# ---------------------------------------------------------------- 실패 기록
def test_scan_exception_is_recorded_as_unreadable_and_logged(
        qapp, isolated_cache, tmp_path, monkeypatch, caplog):
    from aoi_verification.app.coords import scan_image
    e = ex.ExcelExporter(_result(tmp_path), dst_path=tmp_path / "o.xlsx")

    def boom(self, p):
        raise RuntimeError("nas down")
    monkeypatch.setattr(ex.ExcelExporter, "_scan_one", boom)
    p = tmp_path / "a.jpg"
    with caplog.at_level("WARNING", logger="aoi.export"):
        e._scan_prepass([p])
    assert e._scan_cache[str(p)] == (scan_image.UNREADABLE, None, "")
    assert "Scan 확인 실패" in caplog.text


def test_wafer_map_sheet_failure_is_logged_not_silent(
        qapp, isolated_cache, tmp_path, monkeypatch, caplog):
    out = tmp_path / "o.xlsx"
    e = ex.ExcelExporter(_result(tmp_path), dst_path=out,
                         template_path=tmp_path / "none.xlsx",
                         map_renderer=lambda data, px: b"")
    monkeypatch.setattr(ex.ExcelExporter, "_wafer_map_rows", lambda self: ["S1"])

    def boom(self, wb, rows):
        raise RuntimeError("map exploded")
    monkeypatch.setattr(ex.ExcelExporter, "_write_wafer_map_sheet", boom)
    with caplog.at_level("ERROR", logger="aoi.export"):
        e.run()
    assert out.exists()                          # 저장은 계속된다
    assert "map exploded" in caplog.text         # 기록은 남는다


def test_wafer_map_cancel_is_not_swallowed(qapp, isolated_cache, tmp_path, monkeypatch):
    out = tmp_path / "o.xlsx"
    out.write_bytes(b"PREVIOUS")
    e = ex.ExcelExporter(_result(tmp_path), dst_path=out,
                         template_path=tmp_path / "none.xlsx",
                         map_renderer=lambda data, px: b"")
    monkeypatch.setattr(ex.ExcelExporter, "_wafer_map_rows", lambda self: ["S1"])

    def cancelled(self, wb, rows):
        raise ex._Cancelled
    monkeypatch.setattr(ex.ExcelExporter, "_write_wafer_map_sheet", cancelled)
    e.run()
    assert out.read_bytes() == b"PREVIOUS"
