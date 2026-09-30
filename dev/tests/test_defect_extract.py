"""Defect 추출 — 매칭 없이 wafer/LOT 폴더에서 고른 사진과 그 정보를 엑셀로.

사용자 결정(요약):
· 입력: wafer 폴더·LOT 폴더 모두(자동 판별), 한 번에 호기 1개, Camtek·KLA 모두.
· 진입: 설정 화면의 [Defect 추출] — 같은 화면을 한쪽 폴더만 받는 모드로 바꾼다.
· 선별: 매칭의 후보 선별 화면·조작 그대로.
· 엑셀: 양식.xlsx, 시트 1개(세로), A=No · B=Slot · C=사진 · D=정보 — **미매칭 사진
  정보를 적듯이**(D열 글자로 보이게) + Wafer Map 시트.  정렬 slot → 파일명.
· KLA slot명: 매칭과 같다(정보파일 우선, OCR 폴백, KLA 여부는 묻는다).
· 저장: 결과 엑셀과 같은 방식, 이름에 '추출'.  간단한 요약 후 저장.
"""

from __future__ import annotations

import os
from pathlib import Path

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from aoi_verification.app.models.result import (EXTRACT_MODE, FinalResult,  # noqa: E402
                                                MissEntry, extract_result)
from aoi_verification.app.models.slot import (rename_slots_by_wafer_id,   # noqa: E402
                                              scan_one_side)


def _touch_jpeg(p: Path) -> Path:
    p.parent.mkdir(parents=True, exist_ok=True)
    try:
        from PIL import Image
        Image.new("RGB", (40, 30), color=(80, 120, 200)).save(str(p), "JPEG")
    except ImportError:
        p.write_bytes(b"")
    return p


# ── 폴더 판별 ────────────────────────────────────────────────────────────────
def test_lot_folder_gives_one_slot_per_subfolder(tmp_path):
    for slot, names in {"S1": ["b.jpg", "a.jpg"], "S2": ["c.jpg"]}.items():
        for n in names:
            _touch_jpeg(tmp_path / slot / n)
    (tmp_path / "S3").mkdir()                          # 사진 없는 슬롯도 목록엔 있다
    sr = scan_one_side(tmp_path)
    assert sorted(sr.slots) == ["S1", "S2", "S3"]
    assert [it.filename for it in sr.slots["S1"].ref_images] == ["a.jpg", "b.jpg"]
    assert all(it.side == "ref" and it.slot == "S1"
               for it in sr.slots["S1"].ref_images)
    assert sr.slots["S2"].ref_dir == tmp_path / "S2"
    assert not sr.slots["S1"].val_images
    assert sr.ref_only == [] and sr.val_only == []


def test_lot_folder_honours_subset(tmp_path):
    for slot in ("S1", "S2", "S3"):
        _touch_jpeg(tmp_path / slot / "a.jpg")
    assert sorted(scan_one_side(tmp_path, only={"S1", "S3"}).slots) == ["S1", "S3"]


def test_wafer_folder_is_one_slot_even_with_subfolders(tmp_path):
    """★ Camtek wafer 폴더에는 `_recipe_files_…` 같은 하위 폴더가 함께 있다 —
    '하위 폴더가 있으면 LOT' 으로 판정하면 그 폴더들이 슬롯이 되고 사진은 사라진다."""
    wafer = tmp_path / "GIX5703110"
    _touch_jpeg(wafer / "1_2_d.jpg")
    (wafer / "_recipe_files_GX57001305").mkdir()
    sr = scan_one_side(wafer, only={"_recipe_files_GX57001305"})   # only 는 무시된다
    assert list(sr.slots) == ["GIX5703110"]
    assert [it.filename for it in sr.slots["GIX5703110"].ref_images] == ["1_2_d.jpg"]
    assert sr.slots["GIX5703110"].ref_dir == wafer


def test_ignored_photos_do_not_make_a_wafer_folder(tmp_path):
    """웨이퍼 전경 사진만 있는 폴더는 wafer 폴더가 아니다(열거가 걸러 내는 사진)."""
    _touch_jpeg(tmp_path / "CognexInSight17xx_Bottom_Slot21.jpg")
    _touch_jpeg(tmp_path / "S1" / "a.jpg")
    assert list(scan_one_side(tmp_path).slots) == ["S1"]


def test_scan_reports_progress(tmp_path):
    for slot in ("S1", "S2"):
        _touch_jpeg(tmp_path / slot / "a.jpg")
    seen = []
    scan_one_side(tmp_path, progress=lambda d, t: seen.append((d, t)))
    assert seen == [(1, 2), (2, 2)]


# ── KLA slot명 ───────────────────────────────────────────────────────────────
def test_kla_folders_are_renamed_to_wafer_id(tmp_path):
    for slot in ("2026-08-02-23-09_4", "2026-08-03-20-28_2"):
        _touch_jpeg(tmp_path / slot / "x.jpg")
    sr = scan_one_side(tmp_path)
    kla = rename_slots_by_wafer_id(sr, {"2026-08-02-23-09_4": "a1033abqewg3"})
    # 표기는 매칭과 같은 대문자 정규화
    assert sorted(sr.slots) == ["2026-08-03-20-28_2", "A1033ABQEWG3"]
    assert kla == {"A1033ABQEWG3": "2026-08-02-23-09_4"}
    s = sr.slots["A1033ABQEWG3"]
    assert s.name == "A1033ABQEWG3"
    assert all(it.slot == "A1033ABQEWG3" for it in s.ref_images)
    assert s.ref_dir == tmp_path / "2026-08-02-23-09_4"      # 원본 폴더는 그대로


def test_kla_rename_never_overwrites_another_slot(tmp_path):
    """★ 두 폴더가 같은 WaferID 를 내면 두 번째는 폴더명 그대로 남는다 — 덮으면
    한쪽 웨이퍼의 사진이 조용히 사라진다."""
    for slot in ("K1", "K2"):
        _touch_jpeg(tmp_path / slot / "x.jpg")
    sr = scan_one_side(tmp_path)
    kla = rename_slots_by_wafer_id(sr, {"K1": "W1", "K2": "W1"})
    assert sorted(sr.slots) == ["K2", "W1"]
    assert kla == {"W1": "K1"}


# ── 결과 묶음 ────────────────────────────────────────────────────────────────
def test_extract_result_carries_only_picked_photos():
    r = extract_result("3", {"S2": [Path("/x/S2/b.jpg")], "S1": [Path("/x/S1/a.jpg")],
                             "S3": []},
                       kla_folders={"S1": "k"}, slot_numbers={"S1": "6"})
    assert r.mode == EXTRACT_MODE and r.matches == []
    assert [(u.slot, u.path.name, u.side) for u in r.unmatched_refs] == [
        ("S1", "a.jpg", "ref"), ("S2", "b.jpg", "ref")]
    # Wafer map 에는 고른 사진만 — 빈 슬롯은 행을 만들지 않는다.
    assert r.slot_images == {"S1": ([Path("/x/S1/a.jpg")], []),
                             "S2": ([Path("/x/S2/b.jpg")], [])}
    assert r.kla_folders == {"S1": "k"} and r.slot_numbers == {"S1": "6"}
    assert r.val_machine == ""


def test_wafer_map_has_no_match_state_for_extraction(monkeypatch, tmp_path):
    """추출은 매칭을 하지 않았다 — 점을 '미매치' 로 칠하면 거짓이다."""
    from aoi_verification.app.coords import wafer_map as wm
    photo = _touch_jpeg(tmp_path / "S1" / "a.jpg")
    seen = []
    monkeypatch.setattr(wm, "build_map", lambda coords, matched=None:
                        seen.append(matched) or wm.MapData(None, (), ()))
    wm.slot_maps(extract_result("3", {"S1": [photo]}), "S1")
    assert seen == [None, None]
    single = FinalResult(mode="single", ref_machine="1", val_machine="2",
                         slot_images={"S1": ([photo], [])})
    seen.clear()
    wm.slot_maps(single, "S1")
    assert seen == [set(), set()]                     # 매칭 결과는 예전 그대로


# ── 엑셀 ─────────────────────────────────────────────────────────────────────
def _export(result, dst, tpl=None, map_renderer=None):
    from aoi_verification.app.workers.exporter import ExcelExporter
    exp = ExcelExporter(result, dst_path=dst,
                        template_path=tpl or dst.parent / "no_template.xlsx",
                        map_renderer=map_renderer)
    failed = []
    exp.signals.failed.connect(failed.append)
    exp.run()
    assert not failed, failed
    return dst


def _png(_data, size):
    import io

    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (size, size), "white").save(buf, "PNG")
    return buf.getvalue()


def test_excel_layout(qapp, tmp_path):
    pytest.importorskip("openpyxl")
    pytest.importorskip("PIL.Image")
    from openpyxl import load_workbook

    from aoi_verification.app import i18n
    from aoi_verification.app.utils import paths
    from aoi_verification.app.workers import exporter as ex

    a = _touch_jpeg(tmp_path / "src" / "S2" / "a.jpg")
    b = _touch_jpeg(tmp_path / "src" / "S1" / "b.jpg")
    c = _touch_jpeg(tmp_path / "src" / "S1" / "a.jpg")
    result = extract_result("3호기", {"S2": [a], "S1": [b, c]},
                            slot_numbers={"S1": "6"})
    tpl = paths.template_path()
    dst = _export(result, tmp_path / "out.xlsx",
                  tpl=tpl if tpl.exists() else None, map_renderer=_png)

    wb = load_workbook(str(dst), rich_text=True)
    # 시트 1개(요약) + Wafer Map — '미매칭 사진' 시트는 없다.
    assert wb.sheetnames == ["out", i18n.KO.WAFER_MAP_SHEET]
    ws = wb["out"]
    assert ws["C2"].value == "AOI-3"
    assert ws["D2"].value == i18n.KO.EXTRACT_INFO_HEADER
    # 정렬: slot → 파일명
    rows = [(ws[f"A{r}"].value, str(ws[f"B{r}"].value).split("\n")[0],
             str(ws[f"D{r}"].value).split("\n")[0]) for r in (3, 4, 5)]
    assert rows == [(1, "S1", "a.jpg"), (2, "S1", "b.jpg"), (3, "S2", "a.jpg")]
    assert ws["B3"].value == "S1\n(#6)"
    assert len(ws._images) == 3                        # C열 사진
    for r in (3, 4, 5):
        d = ws[f"D{r}"]
        assert d.comment is None                       # '미매칭' 메모 없음
        assert ex.UNMATCHED_FILL not in str(d.fill.fgColor.rgb)   # 미매칭 틴트 없음
    # Wafer Map — 장비가 하나라 맵 칸도 하나.
    wm = wb[i18n.KO.WAFER_MAP_SHEET]
    assert wm["B1"].value.startswith(i18n.KO.EXTRACT_MAP_SHEET_COL)
    assert "AOI-3" in wm["B1"].value
    assert wm["C1"].value is None
    assert [wm.cell(row=r, column=1).value for r in (2, 3, 4)] == [
        i18n.KO.WAFER_MAP_SHEET_ALL, "S1\n(#6)", "S2"]


def test_info_column_is_written_like_an_unmatched_row(qapp, tmp_path):
    """★ 사용자 결정: '마치 미매칭 사진 정보 출력하듯이'.  D열 글자는 매칭 결과의
    미매칭 행이 적는 것과 **한 글자도 다르지 않아야** 한다(같은 생산자)."""
    pytest.importorskip("openpyxl")
    pytest.importorskip("PIL.Image")
    from openpyxl import load_workbook

    p = _touch_jpeg(tmp_path / "src" / "S1" / "a.jpg")
    ext = _export(extract_result("3", {"S1": [p]}), tmp_path / "ext.xlsx")
    ref = _export(FinalResult(mode="single", ref_machine="3", val_machine="4",
                              unmatched_refs=[MissEntry("S1", "ref", p)]),
                  tmp_path / "ref.xlsx")
    d_ext = load_workbook(str(ext), rich_text=True)["ext"]["D3"].value
    d_ref = load_workbook(str(ref), rich_text=True)["ref"]["D3"].value
    assert str(d_ext) == str(d_ref)
    assert "a.jpg" in str(d_ext)


# ── 설정 화면 ────────────────────────────────────────────────────────────────
def test_setup_extract_mode_takes_one_folder(qapp, tmp_path):
    pytest.importorskip("PyQt6.QtWidgets")
    from aoi_verification.app import i18n
    from aoi_verification.app.ui.pages import setup_page as sp

    page = sp.SetupPage()
    try:
        page.ref_path_edit.setText(str(tmp_path))
        page.val_path_edit.setText("")
        assert page._validate() is False               # 매칭은 두 폴더가 필요
        page.set_extract_mode(True)
        assert page.is_extract_mode()
        assert page._validate() is True                # 추출은 한 폴더면 된다
        assert page.val_group.isHidden()
        assert page._setting_cards[1].isHidden()       # 매칭 설정
        assert page._mode_badge_card.isHidden()
        assert not page._setting_cards[0].isHidden()   # 실행 옵션은 남는다
        assert page.start_btn.text() == i18n.KO.BTN_EXTRACT_START
        assert page._title_label.text() == i18n.KO.EXTRACT_TITLE
        assert page.extract_btn.text() == i18n.KO.EXTRACT_BACK_BUTTON
        page.ref_machine_edit.setText("17")
        inp = page._collect_input()
        assert inp.extract is True
        assert inp.ref_root == tmp_path and inp.val_root == tmp_path
        assert inp.ref_machine == "17" and inp.val_machine == ""
        # 되돌리면 원래 화면 그대로(입력값 보존)
        page.set_extract_mode(False)
        assert not page.val_group.isHidden()
        assert not page._setting_cards[1].isHidden()
        assert page.start_btn.text() == i18n.KO.BTN_START
        assert page.ref_path_edit.text() == str(tmp_path)
        assert page._validate() is False
    finally:
        page.deleteLater()


def test_extract_button_sits_in_the_auxiliary_group(qapp):
    """액션바 자리 계약 — 보조 버튼은 stretch 앞, 주 액션은 맨 끝."""
    from aoi_verification.app.ui.pages import setup_page as sp
    page = sp.SetupPage()
    try:
        bar = page._action_bar
        widgets = [bar.itemAt(i).widget() for i in range(bar.count())]
        stretch_at = next(i for i in range(bar.count())
                          if bar.itemAt(i).spacerItem() is not None)
        assert widgets.index(page.extract_btn) < stretch_at
        assert widgets[-1] is page.start_btn
    finally:
        page.deleteLater()


# ── 메인 창 ──────────────────────────────────────────────────────────────────
def test_file_name_follows_result_rule_with_extract(qapp, tmp_path):
    from aoi_verification.app.ui import main_window as mw
    from aoi_verification.app.ui.pages.setup_page import SetupInput

    lot = tmp_path / "LOT"
    (lot / "S1").mkdir(parents=True)
    (lot / "S1" / "WaferInfo.ini").write_text("InputLot=GFW-RDL4\n", encoding="utf-8")
    inp = SetupInput(mode="single", ref_root=lot, val_root=lot, ref_machine="17",
                     val_machine="", threshold=0.7, extract=True)
    assert mw.MainWindow._suggest_result_name(inp) == "17 RDL4_GFW Defect 추출.xlsx"
    inp.ref_root = inp.val_root = tmp_path / "없음"
    assert mw.MainWindow._suggest_result_name(inp) == "AOI 17 Defect 추출.xlsx"


def _pump_until(cond, timeout_s=30.0):
    import time

    from PyQt6.QtWidgets import QApplication
    end = time.monotonic() + timeout_s
    while time.monotonic() < end:
        QApplication.processEvents()
        if cond():
            return True
        time.sleep(0.01)
    return cond()


def test_end_to_end_select_then_save(qapp, tmp_path, monkeypatch):
    """설정(추출) → 스캔 → 썸네일 → 후보 선별(실제 SelectPage) → 요약 → 저장.

    매칭·검토·결과 화면은 거치지 않고, 저장 뒤에는 설정 화면으로 돌아온다."""
    pytest.importorskip("openpyxl")
    pytest.importorskip("PIL.Image")
    from openpyxl import load_workbook

    from aoi_verification.app.ui import main_window as mw
    from aoi_verification.app.ui.pages.setup_page import SetupInput
    from aoi_verification.app.ui.widgets import save_name_dialog

    lot = tmp_path / "LOT"
    for slot, names in {"S1": ["a.jpg", "b.jpg"], "S2": ["c.jpg"]}.items():
        for n in names:
            _touch_jpeg(lot / slot / n)

    monkeypatch.setattr(mw.MainWindow, "_start_backend_import_async",
                        lambda self: None)
    # 개발 모드의 결과 폴더는 저장소 안(`결과/`)이다 — 테스트가 거기에 쓰면 안 된다.
    results = tmp_path / "결과"
    results.mkdir()
    monkeypatch.setattr(mw.paths, "results_dir", lambda: results)
    win = mw.MainWindow()
    win._on_backend_loaded()                  # 나머지 페이지를 지금 만든다
    choices: list = []

    def fake_choose(_parent, title, text, options, **kw):
        keys = [o[0] for o in options]
        choices.append((title, text, keys))
        if "kla" in keys:
            return "none"                     # KLA 아님
        if "single" in keys:
            return "single"                   # 요약 창 — 한 시트로 저장
        return "close"                        # 저장 완료 → 닫기

    monkeypatch.setattr(mw.sheets, "choose", fake_choose)

    class _Dlg:
        def __init__(self, name, _folder, parent=None):
            self.chosen = name

    monkeypatch.setattr(save_name_dialog, "SaveNameDialog", _Dlg)
    monkeypatch.setattr(mw.sheets, "run", lambda dlg, **k: True)
    win.show()                                # 선별은 보이는 화면에서만 결정을 받는다
    try:
        win._on_start(SetupInput(mode="single", ref_root=lot, val_root=lot,
                                 ref_machine="17", val_machine="", threshold=0.7,
                                 extract=True))
        assert _pump_until(lambda: win._stack.currentWidget() is win._select_page)
        # 매칭과 같은 선별 화면 — 3장 중 a, c 를 고르고 b 를 뺀다(방향키와 같은 경로).
        page = win._select_page
        for _ in range(3):
            cur = page._current
            page._decide("exclude" if cur.filename == "b.jpg" else "verify")
            assert _pump_until(lambda: page._current is not cur
                               or not page.get_state().queue, 5)
        assert _pump_until(lambda: win._stack.currentWidget() is win._setup_page)
        summary = [c for c in choices if "single" in c[2]]
        assert summary and "2" in summary[0][1]       # defect 2건
        saved = list(results.glob("*추출*.xlsx"))
        assert len(saved) == 1
        ws = load_workbook(str(saved[0]))[saved[0].stem[:31]]
        names = [str(ws[f"D{r}"].value).split("\n")[0] for r in (3, 4)]
        assert names == ["a.jpg", "c.jpg"]
        # 추출 세션은 이어하기에 남지 않는다
        from aoi_verification.app.models import session as session_mod
        assert session_mod.load() is None
    finally:
        win.close()


# ── Recipe 나누기 ────────────────────────────────────────────────────────────
def _fake_recipes(monkeypatch, table):
    """파일명 → (코드, 이름).  Surface.flt 를 합성하지 않고 분류 규칙만 본다."""
    from aoi_verification.app.workers import exporter as ex
    none = (None, ex.i18n.KO.EXTRACT_RECIPE_NONE)
    monkeypatch.setattr(ex, "recipe_of", lambda p: table.get(Path(p).name, none))


def test_recipe_groups_order_and_none_last(monkeypatch):
    from aoi_verification.app.workers import exporter as ex
    _fake_recipes(monkeypatch, {"a.jpg": (2, "PI"), "b.jpg": (1, "PI_Bubble"),
                                "c.jpg": (2, "PI")})
    got = ex.recipe_groups([Path("/x/a.jpg"), Path("/x/z.jpg"), Path("/x/b.jpg"),
                            Path("/x/c.jpg")])
    assert [(label, [p.name for p in ps]) for label, ps in got] == [
        ("PI_Bubble", ["b.jpg"]), ("PI", ["a.jpg", "c.jpg"]),
        (ex.i18n.KO.EXTRACT_RECIPE_NONE, ["z.jpg"])]


def test_recipe_of_uses_surface_flt_name_then_code(monkeypatch):
    """출처는 Surface.flt recipe(사용자 결정) — 이름이 없으면 코드로 대신한다."""
    from aoi_verification.app.coords import geometry
    from aoi_verification.app.coords.models import DefectGeometry
    from aoi_verification.app.workers import exporter as ex

    def geo(name):
        return geometry.GeometryResult("ok", DefectGeometry(
            1, 1, 1, 1, zone=1, recipe=3, pixel_um=1, recipe_name=name))
    monkeypatch.setattr(geometry, "resolve", lambda p: geo("PI_Bubble"))
    assert ex.recipe_of("/x/a.jpg") == (3, "PI_Bubble")
    monkeypatch.setattr(geometry, "resolve", lambda p: geo(""))
    assert ex.recipe_of("/x/a.jpg") == (3, "Recipe 3")
    monkeypatch.setattr(geometry, "resolve",
                        lambda p: geometry.GeometryResult("no_flt", None))
    assert ex.recipe_of("/x/a.jpg") == (None, ex.i18n.KO.EXTRACT_RECIPE_NONE)


def test_safe_sheet_title():
    from aoi_verification.app.workers.exporter import _safe_sheet_title
    assert _safe_sheet_title("PI/Bubble:[1]", set()) == "PI_Bubble__1_"
    assert _safe_sheet_title("PI", {"PI"}) == "PI (2)"
    assert len(_safe_sheet_title("x" * 40, {"x" * 31})) == 31


def _three_slot_result(tmp_path):
    ps = {n: _touch_jpeg(tmp_path / "src" / slot / n)
          for slot, n in [("S1", "a.jpg"), ("S1", "b.jpg"), ("S1", "c.jpg"),
                          ("S2", "d.jpg")]}
    return extract_result("3", {"S1": [ps["a.jpg"], ps["b.jpg"], ps["c.jpg"]],
                                "S2": [ps["d.jpg"]]})


def test_sheets_layout_one_sheet_per_recipe(qapp, tmp_path, monkeypatch):
    pytest.importorskip("openpyxl")
    pytest.importorskip("PIL.Image")
    from openpyxl import load_workbook

    from aoi_verification.app import i18n
    from aoi_verification.app.workers import exporter as ex

    _fake_recipes(monkeypatch, {"a.jpg": (1, "PI_Bubble"), "c.jpg": (1, "PI_Bubble"),
                                "b.jpg": (2, "PI")})
    dst = tmp_path / "out.xlsx"
    exp = ex.ExcelExporter(_three_slot_result(tmp_path), dst_path=dst,
                           template_path=tmp_path / "none.xlsx", map_renderer=_png,
                           recipe_layout=ex.RECIPE_LAYOUT_SHEETS)
    exp.run()
    wb = load_workbook(str(dst), rich_text=True)
    assert wb.sheetnames == ["PI_Bubble", "PI", i18n.KO.EXTRACT_RECIPE_NONE,
                             i18n.KO.WAFER_MAP_SHEET]

    def rows(name):
        ws = wb[name]
        out, r = [], 3
        while ws[f"A{r}"].value is not None:
            out.append((str(ws[f"B{r}"].value), str(ws[f"D{r}"].value).split("\n")[0]))
            r += 1
        return out
    assert rows("PI_Bubble") == [("S1", "a.jpg"), ("S1", "c.jpg")]
    assert rows("PI") == [("S1", "b.jpg")]
    assert rows(i18n.KO.EXTRACT_RECIPE_NONE) == [("S2", "d.jpg")]
    assert wb["PI"]["D2"].value == i18n.KO.EXTRACT_INFO_HEADER


def test_columns_layout_side_by_side(qapp, tmp_path, monkeypatch):
    """한 시트, Recipe 마다 [사진|정보] 열.  슬롯 안에서는 순서대로 채우고 모자라면 빈칸."""
    pytest.importorskip("openpyxl")
    pytest.importorskip("PIL.Image")
    from openpyxl import load_workbook

    from aoi_verification.app import i18n
    from aoi_verification.app.workers import exporter as ex

    _fake_recipes(monkeypatch, {"a.jpg": (1, "PI_Bubble"), "c.jpg": (1, "PI_Bubble"),
                                "b.jpg": (2, "PI")})
    dst = tmp_path / "out.xlsx"
    single = tmp_path / "single.xlsx"
    ex.ExcelExporter(_three_slot_result(tmp_path), dst_path=dst,
                     template_path=tmp_path / "none.xlsx", map_renderer=_png,
                     recipe_layout=ex.RECIPE_LAYOUT_COLUMNS).run()
    ex.ExcelExporter(_three_slot_result(tmp_path), dst_path=single,
                     template_path=tmp_path / "none.xlsx").run()
    wb = load_workbook(str(dst), rich_text=True)
    assert wb.sheetnames == ["out", i18n.KO.WAFER_MAP_SHEET]
    ws = wb["out"]
    # 머리: Recipe 이름(두 칸 병합) / AOI-N · 정보
    assert [ws[f"{c}1"].value for c in "CEG"] == ["PI_Bubble", "PI",
                                                 i18n.KO.EXTRACT_RECIPE_NONE]
    assert [ws[f"{c}2"].value for c in "CDEF"] == ["AOI-3", i18n.KO.EXTRACT_INFO_HEADER,
                                                  "AOI-3", i18n.KO.EXTRACT_INFO_HEADER]
    assert "C1:D1" in {str(r) for r in ws.merged_cells.ranges}

    def name(cell):
        return None if cell.value is None else str(cell.value).split("\n")[0]
    # S1: PI_Bubble a,c / PI b → 2행 · S2: 없음 d → 1행
    grid = [(ws[f"A{r}"].value, str(ws[f"B{r}"].value), name(ws[f"D{r}"]),
             name(ws[f"F{r}"]), name(ws[f"H{r}"])) for r in (3, 4, 5)]
    assert grid == [(1, "S1", "a.jpg", "b.jpg", None),
                    (2, "S1", "c.jpg", None, None),
                    (3, "S2", None, None, "d.jpg")]
    assert ws["A6"].value is None
    assert len(ws._images) == 4                      # 사진 4장
    # 정보 칸 글자는 한 시트 배치의 D열과 같다(같은 생산자)
    ss = load_workbook(str(single), rich_text=True)["single"]
    assert str(ws["D3"].value) == str(ss["D3"].value)


def test_layout_is_ignored_for_match_results(qapp, tmp_path):
    """Recipe 나누기는 추출 전용 — 매칭 결과에 값을 줘도 예전 배치 그대로다."""
    from aoi_verification.app.workers import exporter as ex
    exp = ex.ExcelExporter(FinalResult(mode="single", ref_machine="1", val_machine="2"),
                           dst_path=tmp_path / "o.xlsx",
                           recipe_layout=ex.RECIPE_LAYOUT_SHEETS)
    assert exp._recipe_layout == ex.RECIPE_LAYOUT_SINGLE


def test_export_logs_stage_times(qapp, tmp_path, caplog):
    """느린 단계를 실측으로 가리기 위한 기록 — app.log 로 가는 'aoi.export' 로거."""
    import logging

    from aoi_verification.app.workers import exporter as ex
    p = _touch_jpeg(tmp_path / "src" / "S1" / "a.jpg")
    with caplog.at_level(logging.INFO, logger="aoi.export"):
        ex.ExcelExporter(extract_result("3", {"S1": [p]}), dst_path=tmp_path / "o.xlsx",
                         template_path=tmp_path / "none.xlsx").run()
    text = caplog.text
    assert "저장 소요 [시트" in text and "저장 소요 [파일 쓰기]" in text


def test_summary_offers_layouts_and_remembers(qapp, tmp_path, monkeypatch):
    """요약 창에서 방식을 고르고, 마지막 선택이 다음 기본이 된다."""
    from aoi_verification.app.ui import main_window as mw
    from aoi_verification.app.ui.pages.setup_page import SetupInput
    from aoi_verification.app.ui.widgets import save_name_dialog
    from aoi_verification.app.utils import prefs
    from aoi_verification.app.workers import exporter as ex

    monkeypatch.setattr(mw.MainWindow, "_start_backend_import_async",
                        lambda self: None)
    win = mw.MainWindow()
    seen, started = [], []
    monkeypatch.setattr(mw.sheets, "choose",
                        lambda *a, **k: (seen.append((a[3], k.get("default"))),
                                         ex.RECIPE_LAYOUT_COLUMNS)[1])

    class _Dlg:
        def __init__(self, name, _folder, parent=None):
            self.chosen = name
    monkeypatch.setattr(save_name_dialog, "SaveNameDialog", _Dlg)
    monkeypatch.setattr(mw.sheets, "run", lambda dlg, **k: True)
    monkeypatch.setattr(mw.MainWindow, "_start_extract_export",
                        lambda self, r, d, layout: started.append(layout))
    try:
        win._input = SetupInput(mode="single", ref_root=tmp_path, val_root=tmp_path,
                                ref_machine="3", val_machine="", threshold=0.7,
                                extract=True)
        win._working_xlsx = tmp_path / "x.xlsx"
        item = mw.ImageItem("S1", _touch_jpeg(tmp_path / "S1" / "a.jpg"), "ref")
        win._finish_extract({"S1": [item]})
        keys = [o[0] for o in seen[0][0]]
        assert keys[:3] == [ex.RECIPE_LAYOUT_SINGLE, ex.RECIPE_LAYOUT_SHEETS,
                            ex.RECIPE_LAYOUT_COLUMNS]
        assert seen[0][1] == ex.RECIPE_LAYOUT_SINGLE          # 처음엔 한 시트
        assert started == [ex.RECIPE_LAYOUT_COLUMNS]
        assert prefs.load().extract_recipe_layout == ex.RECIPE_LAYOUT_COLUMNS
        win._finish_extract({"S1": [item]})
        assert seen[1][1] == ex.RECIPE_LAYOUT_COLUMNS         # 마지막 선택이 기본
    finally:
        win.close()


# ── Recipe 나란히: 같은 결함(150µm 미만)은 같은 행, 위쪽으로 ─────────────────
def _abs(**xy):
    return lambda p: ("abs",) + xy[str(p)] if str(p) in xy else None


def test_pair_rows_puts_close_defects_on_one_row_first():
    from aoi_verification.app.workers.exporter import pair_rows
    x20 = ["a", "b", "c"]
    x5 = ["p", "q"]
    pos = _abs(a=(0.0, 0.0), b=(1000.0, 0.0), c=(5000.0, 0.0),
               p=(5100.0, 0.0),            # c 와 100µm → 같은 행
               q=(1000.0, 149.9))          # b 와 149.9µm → 같은 행
    assert pair_rows([x20, x5], pos) == [["b", "q"], ["c", "p"], ["a", None]]


def test_pair_rows_threshold_is_strict_and_one_to_one():
    from aoi_verification.app.workers.exporter import pair_rows
    # a–p 는 정확히 150µm → '미만' 이 아니므로 짝이 아니다(순서대로 채운 행).
    pos = _abs(a=(0.0, 0.0), b=(1000.0, 0.0), p=(150.0, 0.0), q=(1003.0, 0.0))
    assert pair_rows([["a", "b"], ["p", "q"]], pos) == [["b", "q"], ["a", "p"]]
    # q 는 a(2µm)·b(3µm) 둘 다 가깝지만 한 행에만 — 더 가까운 a 와.
    pos = _abs(a=(0.0, 0.0), b=(5.0, 0.0), q=(2.0, 0.0))
    assert pair_rows([["a", "b"], ["q"]], pos) == [["a", "q"], ["b", None]]


def test_pair_rows_three_recipes_needs_all_close():
    from aoi_verification.app.workers.exporter import pair_rows
    pos = _abs(a=(0.0, 0.0), p=(100.0, 0.0), u=(200.0, 0.0), v=(50.0, 0.0))
    # u 는 p 와 100 이지만 a 와 200 → a·p 행에 못 들어간다.  v 는 둘 다 150 미만.
    assert pair_rows([["a"], ["p"], ["u", "v"]], pos) == [["a", "p", "v"],
                                                          [None, None, "u"]]


def test_pair_rows_die_coords_compare_only_within_same_die():
    from aoi_verification.app.workers.exporter import pair_rows
    pos = {"a": ("die", 1, 2, 10.0, 10.0), "p": ("die", 1, 2, 20.0, 10.0),
           "b": ("die", 1, 3, 10.0, 10.0), "q": ("die", 1, 4, 10.0, 10.0)}.get
    assert pair_rows([["a", "b"], ["q", "p"]], pos) == [["a", "p"], ["b", "q"]]


def test_columns_sheet_pairs_rows_on_top(qapp, tmp_path, monkeypatch):
    pytest.importorskip("openpyxl")
    pytest.importorskip("PIL.Image")
    from openpyxl import load_workbook

    from aoi_verification.app.workers import exporter as ex

    _fake_recipes(monkeypatch, {"a.jpg": (1, "x20"), "b.jpg": (1, "x20"),
                                "c.jpg": (2, "x5"), "d.jpg": (2, "x5")})
    xy = {"a.jpg": (0.0, 0.0), "b.jpg": (9000.0, 0.0),
          "c.jpg": (50000.0, 0.0), "d.jpg": (9050.0, 30.0)}   # b–d 58µm
    monkeypatch.setattr(ex, "defect_position",
                        lambda p: ("abs",) + xy[Path(p).name])
    ps = [_touch_jpeg(tmp_path / "src" / "S1" / n)
          for n in ("a.jpg", "b.jpg", "c.jpg", "d.jpg")]
    dst = tmp_path / "out.xlsx"
    ex.ExcelExporter(extract_result("3", {"S1": ps}), dst_path=dst,
                     template_path=tmp_path / "none.xlsx",
                     recipe_layout=ex.RECIPE_LAYOUT_COLUMNS).run()
    ws = load_workbook(str(dst), rich_text=True)["out"]

    def name(cell):
        return None if cell.value is None else str(cell.value).split("\n")[0]
    assert [(name(ws[f"D{r}"]), name(ws[f"F{r}"])) for r in (3, 4)] == [
        ("b.jpg", "d.jpg"), ("a.jpg", "c.jpg")]
    # 따로 표시하지 않는다 — 두 행의 정보 칸 글꼴·채움이 같은 규칙(줄무늬)을 따른다.
    assert ws["D3"].comment is None and ws["D4"].comment is None


def test_choice_buttons_wrap_instead_of_clipping(styled_qapp):
    """버튼이 시트 폭에 한 줄로 안 들어가면 두 열로 접는다 — 글자가 잘리지 않는다."""
    from PyQt6.QtWidgets import QGridLayout

    from aoi_verification.app import i18n
    from aoi_verification.app.ui.widgets import sheet_host as sh
    K = i18n.KO
    many = [("s", K.EXTRACT_SAVE_SINGLE, "primary"),
            ("h", K.EXTRACT_SAVE_SHEETS, "ghost"),
            ("c", K.EXTRACT_SAVE_COLUMNS, "ghost"),
            ("b", K.EXTRACT_SUMMARY_BACK, "ghost")]
    d = sh._ChoiceSheet("t", "body", many, default="s")
    try:
        d.show()
        for _ in range(3):
            styled_qapp.processEvents()
        assert d.width() <= d.maximumWidth()
        for b in d._buttons.values():
            assert b.width() >= b.sizeHint().width(), b.text()
        grids = [d.layout().itemAt(i).layout() for i in range(d.layout().count())]
        assert any(isinstance(g, QGridLayout) for g in grids)
    finally:
        d.deleteLater()
    few = sh._ChoiceSheet("t", "body", many[:2], default="s")
    try:
        grids = [few.layout().itemAt(i).layout() for i in range(few.layout().count())]
        assert not any(isinstance(g, QGridLayout) for g in grids)   # 들어가면 예전 그대로
    finally:
        few.deleteLater()
