"""결과 엑셀의 Scan image 열 — 실제 `wb.save` 후 다시 열어 검증(합성 이미지).

사용자 결정:
- 이번 저장의 사진 중 하나라도 Scan 이 확인되면 열을 넣는다(묶어서 판단).  없으면 예전 그대로.
- 기존 열 뒤에 붙인다 — 매칭 C 기준·D 검증·E 기준 Scan·F 검증 Scan.  미매칭 행은 D 정보,
  E 그 사진의 Scan, F 빈칸.
- Scan 열이 있는 시트에서 그 행만 Scan 이 없으면 짧은 문구.
- 추출은 C 사진·D 정보·E Scan, Recipe 나란히 배치는 [사진|정보|Scan].
- (2026-10-26) Scan 칸에도 Scan 파일명을 셀 값으로 적는다.  Scan 목록이 없는 장비 쪽(KLA)은
  열을 만들지 않는다 — 그러면 Camtek 쪽 Scan 하나만 E 에 온다.
"""
from __future__ import annotations

import io
import os
import zipfile
from pathlib import Path

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
pytest.importorskip("PyQt6.QtWidgets")
pytest.importorskip("openpyxl")
pytest.importorskip("PIL.Image")

from openpyxl import load_workbook                              # noqa: E402
from PIL import Image                                           # noqa: E402

from aoi_verification.app import i18n                           # noqa: E402
from aoi_verification.app.coords import scan_image              # noqa: E402
from aoi_verification.app.models.result import (                # noqa: E402
    FinalResult, MatchResult, MissEntry, extract_result)
from aoi_verification.app.utils import paths                    # noqa: E402
from aoi_verification.app.workers import exporter as ex         # noqa: E402

SCAN_NAME = "-1.2.t.jpeg"


@pytest.fixture(autouse=True)
def _fresh():
    scan_image.clear_caches()
    yield
    scan_image.clear_caches()


def _folder(tmp_path, name, *, scan=True):
    """웨이퍼 폴더 — 픽셀 1.0㎛, Scan 1장(640×400, 중심 10000/20000 → X 9680~10320)."""
    f = tmp_path / name
    f.mkdir()
    (f / "Params_WaferInfo.ini").write_text("[I]\nRefPixelSizeX=1.0\nRefPixelSizeY=1.0\n")
    if scan:
        Image.new("L", (640, 400), 100).save(f / SCAN_NAME)
        (f / scan_image.SCAN_LIST_NAME).write_text(f"Version=1\n{SCAN_NAME},10000,20000\n")
    return f


def _color(folder, x, y, tag="1"):
    p = folder / f"{x}.{y}.c.{tag}.jpeg"
    Image.new("RGB", (138, 104), (40, 90, 60)).save(p)
    return p


def _anchors(ws) -> set[tuple[str, int]]:
    from openpyxl.utils import get_column_letter
    return {(get_column_letter(im.anchor._from.col + 1), im.anchor._from.row + 1)
            for im in ws._images}


def _match_result(tmp_path):
    ref_f = _folder(tmp_path, "ref")
    val_f = _folder(tmp_path, "val")
    a_ref = _color(ref_f, 10000, 20000)          # Scan 안
    a_val = _color(val_f, 99999, 20000)          # Scan 밖 → 'Scan 없음'
    b_ref = _color(ref_f, 10100, 20050, "2")     # 미매칭, Scan 안
    return FinalResult(
        mode="single", ref_machine="17", val_machine="23",
        matches=[MatchResult(slot="S1", ref_path=a_ref, val_path=a_val, score=0.9)],
        unmatched_refs=[MissEntry(slot="S1", side="ref", path=b_ref, note="")])


def _export(result, dst, **kw):
    ex.ExcelExporter(result, dst_path=dst, **kw).run()
    assert dst.exists()
    return load_workbook(str(dst), rich_text=True)


def test_match_summary_gets_scan_columns_after_existing(qapp, isolated_cache, tmp_path):
    wb = _export(_match_result(tmp_path), tmp_path / "out.xlsx",
                 template_path=tmp_path / "none.xlsx")
    ws = wb["out"]
    assert [ws[f"{c}2"].value for c in "CDEF"] == ["AOI-17", "AOI-23",
                                                   "AOI-17 Scan", "AOI-23 Scan"]
    assert "C1:F1" in {str(r) for r in ws.merged_cells.ranges}
    a = _anchors(ws)
    # 매치 행(3): C·D 사진, E 기준 Scan.  검증 사진은 Scan 밖이라 F 는 문구.
    assert {("C", 3), ("D", 3), ("E", 3)} <= a and ("F", 3) not in a
    assert ws["F3"].value == i18n.KO.SCAN_EXCEL_NONE
    # 미매칭 행(4): C 사진, D 정보 그대로, E 그 사진의 Scan, F 빈칸.
    assert ("C", 4) in a and ("E", 4) in a and ("D", 4) not in a
    assert "10100.20050.c.2.jpeg" in str(ws["D4"].value)
    assert ws["F4"].value is None
    assert ws.column_dimensions["E"].width == ws.column_dimensions["C"].width
    # 미매칭 시트도 같은 열 구성
    um = wb[i18n.KO.SHEET_UNMATCHED]
    assert um["E2"].value == "AOI-17 Scan" and ("E", 3) in _anchors(um)


def test_no_scan_anywhere_keeps_old_layout(qapp, isolated_cache, tmp_path):
    f = _folder(tmp_path, "w", scan=False)
    a, b = _color(f, 1, 2), _color(f, 3, 4, "2")
    result = FinalResult(mode="single", ref_machine="1", val_machine="2",
                         matches=[MatchResult(slot="S1", ref_path=a, val_path=b, score=1)],
                         unmatched_refs=[])
    ws = _export(result, tmp_path / "out.xlsx", template_path=tmp_path / "n.xlsx")["out"]
    assert ws.max_column == 4
    assert _anchors(ws) == {("C", 3), ("D", 3)}


def test_embedded_scan_is_the_shared_crop(qapp, isolated_cache, tmp_path):
    dst = tmp_path / "out.xlsx"
    _export(_match_result(tmp_path), dst, template_path=tmp_path / "none.xlsx")
    sizes = set()
    with zipfile.ZipFile(dst) as z:
        for n in z.namelist():
            if n.startswith("xl/media/"):
                sizes.add(Image.open(io.BytesIO(z.read(n))).size)
    assert (300, 300) in sizes          # 300㎛ ÷ 1.0㎛/px — 원본 전체(640×400)가 아니다
    assert (640, 400) not in sizes


def test_unreadable_scan_writes_message(qapp, isolated_cache, tmp_path, monkeypatch):
    monkeypatch.setattr(scan_image, "load_crop", lambda m: None)
    ws = _export(_match_result(tmp_path), tmp_path / "out.xlsx",
                 template_path=tmp_path / "none.xlsx")["out"]
    assert ws["E3"].value.startswith(i18n.KO.SCAN_EXCEL_UNREADABLE)
    assert ws["E2"].value == "AOI-17 Scan"       # 판정은 통과했으니 열은 있다


def test_extract_single_puts_scan_after_info(qapp, isolated_cache, tmp_path):
    f = _folder(tmp_path, "w")
    p = _color(f, 10000, 20000)
    ws = _export(extract_result("3", {"S1": [p]}), tmp_path / "e.xlsx",
                 template_path=tmp_path / "none.xlsx")["e"]
    assert [ws[f"{c}2"].value for c in "CDE"] == ["AOI-3", i18n.KO.EXTRACT_INFO_HEADER,
                                                  i18n.KO.SCAN_EXCEL_HEADER]
    assert ws["F2"].value is None
    assert _anchors(ws) == {("C", 3), ("E", 3)}
    assert p.name in str(ws["D3"].value)


def test_extract_columns_layout_uses_three_column_stride(qapp, isolated_cache, tmp_path,
                                                         monkeypatch):
    f = _folder(tmp_path, "w")
    a, b = _color(f, 10000, 20000, "1"), _color(f, 10050, 20010, "2")
    table = {a.name: (1, "PI"), b.name: (2, "PI2")}
    monkeypatch.setattr(ex, "recipe_of", lambda q: table[Path(q).name])
    ws = _export(extract_result("3", {"S1": [a, b]}), tmp_path / "c.xlsx",
                 template_path=tmp_path / "none.xlsx",
                 recipe_layout=ex.RECIPE_LAYOUT_COLUMNS)["c"]
    assert [ws[f"{c}1"].value for c in "CF"] == ["PI", "PI2"]
    assert {"C1:E1", "F1:H1"} <= {str(r) for r in ws.merged_cells.ranges}
    assert [ws[f"{c}2"].value for c in "DEGH"] == [
        i18n.KO.EXTRACT_INFO_HEADER, i18n.KO.SCAN_EXCEL_HEADER,
        i18n.KO.EXTRACT_INFO_HEADER, i18n.KO.SCAN_EXCEL_HEADER]
    assert {("C", 3), ("E", 3), ("F", 3), ("H", 3)} <= _anchors(ws)


# --- Scan 파일명 · 장비 쪽별 열 (사용자 요청 2026-10-26) -----------------------------
def test_scan_cell_holds_scan_file_name(qapp, isolated_cache, tmp_path):
    """Color 사진 칸처럼 Scan 칸에도 파일명을 셀 값으로 적는다(그림이 위를 덮는다)."""
    ws = _export(_match_result(tmp_path), tmp_path / "out.xlsx",
                 template_path=tmp_path / "none.xlsx")["out"]
    assert ws["E3"].value == SCAN_NAME and ("E", 3) in _anchors(ws)
    assert ws["E4"].value == SCAN_NAME                       # 미매칭 행도
    assert ws["F3"].value == i18n.KO.SCAN_EXCEL_NONE         # 후보가 없으면 이름도 없다


def test_unreadable_scan_message_keeps_file_name(qapp, isolated_cache, tmp_path,
                                                 monkeypatch):
    monkeypatch.setattr(scan_image, "load_crop", lambda m: None)
    ws = _export(_match_result(tmp_path), tmp_path / "out.xlsx",
                 template_path=tmp_path / "none.xlsx")["out"]
    assert ws["E3"].value == f"{i18n.KO.SCAN_EXCEL_UNREADABLE}\n{SCAN_NAME}"


def _kla_folder(tmp_path, name):
    """KLA 처럼 Scan 목록이 아예 없는 폴더."""
    f = tmp_path / name
    f.mkdir()
    return f


@pytest.mark.parametrize("kla_side", ["ref", "val"])
def test_kla_side_gets_no_scan_column(qapp, isolated_cache, tmp_path, kla_side):
    cam = _folder(tmp_path, "cam")
    kla = _kla_folder(tmp_path, "kla")
    cam_a = _color(cam, 10000, 20000)
    kla_a = _color(kla, 10000, 20000)
    cam_u = _color(cam, 10100, 20050, "2")
    kla_u = _color(kla, 5, 6, "2")
    ref_a, val_a = (kla_a, cam_a) if kla_side == "ref" else (cam_a, kla_a)
    result = FinalResult(
        mode="single", ref_machine="17", val_machine="23",
        matches=[MatchResult(slot="S1", ref_path=ref_a, val_path=val_a, score=0.9)],
        unmatched_refs=[
            MissEntry(slot="S1", side=kla_side, path=kla_u, note=""),
            MissEntry(slot="S1", side="val" if kla_side == "ref" else "ref",
                      path=cam_u, note="")])
    wb = _export(result, tmp_path / "out.xlsx", template_path=tmp_path / "none.xlsx")
    ws = wb["out"]
    cam_machine = "AOI-23" if kla_side == "ref" else "AOI-17"
    # Scan 열은 Camtek 쪽 하나뿐 — E 하나, F 는 없다.
    assert ws["E2"].value == f"{cam_machine} Scan"
    assert ws.max_column == 5
    assert "C1:E1" in {str(r) for r in ws.merged_cells.ranges}
    a = _anchors(ws)
    assert ("E", 3) in a and ws["E3"].value == SCAN_NAME     # 매치 행 = Camtek 사진의 Scan
    rows = {ws[f"E{r}"].value for r in (4, 5)}
    assert SCAN_NAME in rows and None in rows                # KLA 미매칭 행은 빈칸
    um = wb[i18n.KO.SHEET_UNMATCHED]                         # 미매칭 시트도 같은 구성
    assert um["E2"].value == f"{cam_machine} Scan" and um.max_column == 5
