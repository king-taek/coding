"""GL 분석 결과 → 엑셀 한 장(통합 문서).  시트: 분석 조건 · ROI 통계 · 히스토그램 · 선 프로파일.

이 모듈은 값을 **그대로 옮기기만** 한다(계산은 ``analysis``).  엑셀 저장 관례는 결과 엑셀과
같다 — 임시 파일에 쓰고, 회사 메타데이터를 걷어 낸 뒤 ``os.replace`` 로 교체한다.
"""

from __future__ import annotations

import os
from pathlib import Path
from typing import Optional, Sequence

import numpy as np

MAX_SHEET_ROWS = 1_048_576


def _stat_cells(row: dict) -> list:
    n = row["n"]
    pct = lambda c: (c / n) if n else None  # noqa: E731
    return [row["name"], row.get("box", ""), n, row["mean"], row["std"], row["min"],
            row["max"], row["below"], pct(row["below"]), row["above"], pct(row["above"])]


def build_workbook(info: Sequence[tuple[str, object]], stats_rows: Sequence[dict],
                   hist_cols: Sequence[tuple[str, np.ndarray]],
                   profiles: Sequence[tuple[str, np.ndarray]], titles: dict):
    """통합 문서를 만든다.

    ``hist_cols`` — [(이름, 길이 2**bits 의 히스토그램)].  모든 열이 0 인 GL 행은 뺀다.
    ``profiles``  — [(이름, 위치별 GL 배열)].  ``titles`` — 시트·열 이름(한국어, ko.py)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    bold = Font(bold=True)

    ws = wb.active
    ws.title = titles["info"]
    for k, v in info:
        ws.append([k, v])
    for c in ws["A"]:
        c.font = bold
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 48

    ws = wb.create_sheet(titles["stats"])
    head = titles["stats_cols"]
    ws.append(head)
    for c in ws[1]:
        c.font = bold
    for r in stats_rows:
        ws.append(_stat_cells(r))
    for i, w in enumerate((16, 22, 12, 12, 12, 10, 10, 12, 10, 12, 10), start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for idx in (3, 4):
            if row[idx].value is not None:
                row[idx].number_format = "0.00"
        for idx in (8, 10):
            row[idx].number_format = "0.00%"

    ws = wb.create_sheet(titles["hist"])
    ws.append([titles["gl"]] + [n for n, _h in hist_cols])
    for c in ws[1]:
        c.font = bold
    if hist_cols:
        stack = np.vstack([np.asarray(h, dtype=np.int64) for _n, h in hist_cols])
        for gl in np.flatnonzero(stack.sum(axis=0))[:MAX_SHEET_ROWS - 1]:
            ws.append([int(gl)] + [int(v) for v in stack[:, gl]])

    ws = wb.create_sheet(titles["profile"])
    ws.append([titles["pos"]] + [n for n, _v in profiles])
    for c in ws[1]:
        c.font = bold
    longest = max((len(v) for _n, v in profiles), default=0)
    for i in range(min(longest, MAX_SHEET_ROWS - 1)):
        ws.append([i] + [int(v[i]) if i < len(v) else None for _n, v in profiles])
    return wb


def save_workbook(wb, dst) -> None:
    """임시 파일에 쓰고 메타데이터를 걷은 뒤 교체한다 — 실패해도 기존 파일은 그대로다."""
    dst = Path(dst)
    tmp = dst.with_name(dst.name + ".part")
    try:
        wb.save(str(tmp))
        try:
            from ..workers.exporter import _strip_corporate_metadata
            _strip_corporate_metadata(tmp)
        except Exception:
            pass          # 메타데이터 정리는 가능할 때만(Qt 없는 환경 · 정리 실패는 치명적이지 않다)
        os.replace(tmp, dst)
    finally:
        if tmp.exists():
            try:
                tmp.unlink()
            except OSError:
                pass
